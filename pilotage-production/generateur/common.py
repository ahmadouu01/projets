# -*- coding: utf-8 -*-
"""Socle de style et helpers communs au classeur Pilotage Production V2."""
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side, NamedStyle
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.properties import PageSetupProperties
from openpyxl.formatting.rule import FormulaRule, ColorScaleRule, CellIsRule

FONT = "Arial"

# --- Palette industrielle ---------------------------------------------------
NAVY      = "0F2A44"   # bandeaux de titre
STEEL     = "1F4E79"   # en-tetes de tableau
SLATE     = "35566F"   # sous-en-tetes
LIGHT     = "EDF2F7"   # bandes de section
ROWALT    = "F5F8FB"   # lignes alternees
BORDERCOL = "B7C4D1"
WHITE     = "FFFFFF"
INPUT_BG  = "FFF9DB"   # cellules a saisir
CALC_BG   = "EAF1F8"   # cellules calculees
REF_BG    = "F2F7ED"   # cellules de referentiel
GREEN     = "1E7B34"
GREEN_BG  = "DFF3E3"
AMBER     = "9C6500"
AMBER_BG  = "FFF0C7"
RED       = "B01919"
RED_BG    = "FBE0E0"
GREY      = "6B7A88"

THIN = Side(style="thin", color=BORDERCOL)
MED  = Side(style="medium", color=STEEL)
BOX  = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)


def f(size=10, bold=False, color="1A1A1A", italic=False):
    return Font(name=FONT, size=size, bold=bold, color=color, italic=italic)


def fill(hexcode):
    return PatternFill("solid", fgColor=hexcode)


def title_band(ws, row, last_col, text, sub=None, height=30):
    """Bandeau de titre principal."""
    ws.merge_cells(start_row=row, start_column=1, end_row=row, end_column=last_col)
    c = ws.cell(row=row, column=1, value=text)
    c.font = Font(name=FONT, size=15, bold=True, color=WHITE)
    c.fill = fill(NAVY)
    c.alignment = Alignment(horizontal="left", vertical="center", indent=1)
    ws.row_dimensions[row].height = height
    for col in range(1, last_col + 1):
        ws.cell(row=row, column=col).fill = fill(NAVY)
    if sub is not None:
        ws.merge_cells(start_row=row + 1, start_column=1, end_row=row + 1, end_column=last_col)
        c2 = ws.cell(row=row + 1, column=1, value=sub)
        c2.font = Font(name=FONT, size=9, color="4A5A6A", italic=True)
        c2.alignment = Alignment(horizontal="left", vertical="center", indent=1)
        ws.row_dimensions[row + 1].height = 18


def section(ws, row, last_col, text, tone=LIGHT, color=STEEL):
    ws.merge_cells(start_row=row, start_column=1, end_row=row, end_column=last_col)
    c = ws.cell(row=row, column=1, value=text)
    c.font = Font(name=FONT, size=10.5, bold=True, color=color)
    c.alignment = Alignment(horizontal="left", vertical="center", indent=1)
    for col in range(1, last_col + 1):
        ws.cell(row=row, column=col).fill = fill(tone)
        ws.cell(row=row, column=col).border = Border(bottom=Side(style="thin", color=color))
    ws.row_dimensions[row].height = 21


def header_row(ws, row, headers, start_col=1, bg=STEEL, height=32, wrap=True):
    """Ecrit une ligne d'en-tetes de tableau."""
    for i, h in enumerate(headers):
        c = ws.cell(row=row, column=start_col + i, value=h)
        c.font = Font(name=FONT, size=9, bold=True, color=WHITE)
        c.fill = fill(bg)
        c.alignment = Alignment(horizontal="center", vertical="center", wrap_text=wrap)
        c.border = BOX
    ws.row_dimensions[row].height = height


def label(ws, ref, text, bold=True, size=10, color="1A1A1A", align="left"):
    c = ws[ref]
    c.value = text
    c.font = f(size, bold, color)
    c.alignment = Alignment(horizontal=align, vertical="center")
    return c


def widths(ws, mapping):
    for col, w in mapping.items():
        ws.column_dimensions[col].width = w


def apply_range(ws, ref, font=None, fillc=None, numfmt=None, align=None,
                border=BOX, wrap=None):
    for row in ws[ref]:
        for c in row:
            if font is not None:
                c.font = font
            if fillc is not None:
                c.fill = fill(fillc)
            if numfmt is not None:
                c.number_format = numfmt
            if align is not None:
                c.alignment = Alignment(horizontal=align, vertical="center",
                                        wrap_text=bool(wrap))
            if border is not None:
                c.border = border


def zebra(ws, first_row, last_row, first_col, last_col):
    for r in range(first_row, last_row + 1):
        if (r - first_row) % 2 == 1:
            for c in range(first_col, last_col + 1):
                cell = ws.cell(row=r, column=c)
                if cell.fill is None or cell.fill.fgColor.rgb in (None, "00000000"):
                    cell.fill = fill(ROWALT)


def page(ws, orientation="landscape", fit_width=1, title_rows=None, area=None):
    ws.page_setup.orientation = orientation
    ws.page_setup.paperSize = ws.PAPERSIZE_A4
    ws.sheet_properties.pageSetUpPr = PageSetupProperties(fitToPage=True)
    ws.page_setup.fitToWidth = fit_width
    ws.page_setup.fitToHeight = 0
    ws.print_options.horizontalCentered = True
    ws.page_margins.left = ws.page_margins.right = 0.4
    ws.page_margins.top = 0.5
    ws.page_margins.bottom = 0.5
    if title_rows:
        ws.print_title_rows = title_rows
    if area:
        ws.print_area = area
    ws.oddFooter.left.text = "&F  —  &A"
    ws.oddFooter.left.size = 8
    ws.oddFooter.right.text = "Page &P / &N"
    ws.oddFooter.right.size = 8


def note(ws, ref, text, span=None):
    """Note methodologique en bas de tableau."""
    c = ws[ref]
    c.value = text
    c.font = Font(name=FONT, size=8.5, italic=True, color=GREY)
    c.alignment = Alignment(horizontal="left", vertical="top", wrap_text=True)
    if span:
        ws.merge_cells(span)
    return c


# Formats numeriques
PCT   = "0.0%"
PCT0  = "0%"
NUM   = "#,##0"
NUM1  = "#,##0.0"
MIN   = '#,##0" min"'
DAYS  = '#,##0" j"'
PPM   = '#,##0" ppm"'
EUR   = '#,##0" €"'
DATE  = "dd/mm/yyyy"
HOUR  = "hh:mm"


def band(ws, row, c1, c2, text, tone=LIGHT, color=STEEL, size=10.5):
    """Bandeau de section sur une plage de colonnes."""
    ws.merge_cells(start_row=row, start_column=c1, end_row=row, end_column=c2)
    c = ws.cell(row=row, column=c1, value=text)
    c.font = Font(name=FONT, size=size, bold=True, color=color)
    c.alignment = Alignment(horizontal="left", vertical="center", indent=1)
    for col in range(c1, c2 + 1):
        ws.cell(row=row, column=col).fill = fill(tone)
        ws.cell(row=row, column=col).border = Border(bottom=Side(style="thin", color=color))
    ws.row_dimensions[row].height = 20


def kpi_tile(ws, row, c1, c2, titre, formule, numfmt, sous1=None, sous2=None):
    """Carte indicateur : titre / valeur / 2 lignes de contexte."""
    ws.merge_cells(start_row=row, start_column=c1, end_row=row, end_column=c2)
    t = ws.cell(row=row, column=c1, value=titre)
    t.font = Font(name=FONT, size=8.5, bold=True, color="FFFFFF")
    t.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
    for col in range(c1, c2 + 1):
        ws.cell(row=row, column=col).fill = fill(STEEL)
    ws.row_dimensions[row].height = 24

    ws.merge_cells(start_row=row + 1, start_column=c1, end_row=row + 2, end_column=c2)
    v = ws.cell(row=row + 1, column=c1, value=formule)
    v.font = Font(name=FONT, size=20, bold=True, color="0F2A44")
    v.alignment = Alignment(horizontal="center", vertical="center")
    v.number_format = numfmt
    ws.row_dimensions[row + 1].height = 20
    ws.row_dimensions[row + 2].height = 16

    for i, txt in enumerate((sous1, sous2)):
        r = row + 3 + i
        ws.merge_cells(start_row=r, start_column=c1, end_row=r, end_column=c2)
        c = ws.cell(row=r, column=c1, value=txt)
        c.font = Font(name=FONT, size=8, color="4A5A6A")
        c.alignment = Alignment(horizontal="center", vertical="center")
        ws.row_dimensions[r].height = 14
    # encadrement
    for r in range(row, row + 5):
        for col in range(c1, c2 + 1):
            cell = ws.cell(row=r, column=col)
            lft = Side(style="thin", color=STEEL) if col == c1 else None
            rgt = Side(style="thin", color=STEEL) if col == c2 else None
            top = Side(style="thin", color=STEEL) if r == row else None
            bot = Side(style="thin", color=STEEL) if r == row + 4 else None
            cell.border = Border(left=lft, right=rgt, top=top, bottom=bot)
            if r > row and cell.fill.fgColor.rgb in (None, "00000000"):
                cell.fill = fill(WHITE)


def str_categories(chart, ref):
    """Force des categories textuelles (strRef) plutot que numeriques."""
    from openpyxl.chart.data_source import AxDataSource, StrRef
    for s in chart.series:
        s.cat = AxDataSource(strRef=StrRef(f=ref))
