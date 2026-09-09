# -*- coding: utf-8 -*-
"""Onglet ARRETS : journal evenementiel des arrets (niveau de detail Pareto)."""
from openpyxl.styles import Alignment, Border, Side
from openpyxl.worksheet.datavalidation import DataValidation
from openpyxl.formatting.rule import FormulaRule
from common import *
import refs as R
import data as D

FIRST, LAST = 5, 404

COLS = [
    ("A", "Date", 11, "in", DATE),
    ("B", "Équipe", 20, "in", None),
    ("C", "Ligne / ressource", 22, "in", None),
    ("D", "Machine / poste", 18, "in", None),
    ("E", "Heure de début", 12, "in", HOUR),
    ("F", "Heure de fin", 12, "in", HOUR),
    ("G", "Durée\n(min)", 9, "ca", NUM),
    ("H", "Code cause", 11, "in", None),
    ("I", "Libellé de la cause", 34, "ca", None),
    ("J", "Famille 5M", 14, "ca", None),
    ("K", "Type", 11, "ca", None),
    ("L", "Rubrique de perte", 22, "ca", None),
    ("M", "Constat terrain", 34, "in", None),
    ("N", "Action immédiate appliquée", 34, "in", None),
    ("O", "N° d'action liée", 13, "in", None),
    ("P", "Niveau d'alerte", 26, "ca", None),
    ("Q", "Filtre cockpit", 10, "ca", "0"),
]

FORMULES = {
    "G": '=IF(OR($E{r}="",$F{r}=""),"",ROUND(($F{r}-$E{r}+IF($F{r}<$E{r},1,0))*1440,0))',
    "I": '=IF($H{r}="","",IFERROR(INDEX(PARAMETRES!$S$7:$S$46,MATCH($H{r},PARAMETRES!$R$7:$R$46,0)),"Code inconnu"))',
    "J": '=IF($H{r}="","",IFERROR(INDEX(PARAMETRES!$T$7:$T$46,MATCH($H{r},PARAMETRES!$R$7:$R$46,0)),""))',
    "K": '=IF($H{r}="","",IFERROR(INDEX(PARAMETRES!$U$7:$U$46,MATCH($H{r},PARAMETRES!$R$7:$R$46,0)),""))',
    "L": '=IF($H{r}="","",IFERROR(INDEX(PARAMETRES!$V$7:$V$46,MATCH($H{r},PARAMETRES!$R$7:$R$46,0)),""))',
    "P": ('=IF($G{r}="","",IF($G{r}>=60,"Arrêt majeur — analyse 5 pourquoi",'
          'IF($G{r}>=30,"À analyser en revue quotidienne",'
          'IF($K{r}="Planifié","Arrêt planifié","Aléa courant"))))'),
    "Q": ('=IF($A{r}="",0,IF(AND(OR({FL}="Toutes",$C{r}={FL}),'
          'OR({FE}="Toutes",$B{r}={FE})),1,0))'),
}


def build(wb, arrets):
    ws = wb.create_sheet("ARRETS")
    ws.sheet_view.showGridLines = False
    ws.sheet_properties.tabColor = "2E86AB"
    n = len(COLS)

    title_band(ws, 1, n, "JOURNAL DES ARRÊTS  —  niveau événementiel",
               "Un enregistrement par arrêt. Ce niveau de détail alimente l'analyse de Pareto par cause et par famille 5M. "
               "Les colonnes bleues sont renseignées automatiquement à partir du code cause.")

    header_row(ws, 4, [c[1] for c in COLS], height=38)
    for i, (lt, _, w, kind, fmt) in enumerate(COLS):
        ws.column_dimensions[lt].width = w
        if kind != "in":
            ws.cell(row=4, column=i + 1).fill = fill(NAVY)
        rng = "%s%d:%s%d" % (lt, FIRST, lt, LAST)
        if kind == "in":
            apply_range(ws, rng, font=f(9, color="0000FF"), fillc=INPUT_BG, numfmt=fmt,
                        align="center" if fmt else "left")
        else:
            apply_range(ws, rng, font=f(9), fillc=CALC_BG, numfmt=fmt,
                        align="center" if fmt else "left")

    for r in range(FIRST, LAST + 1):
        for lt, formule in FORMULES.items():
            ws["%s%d" % (lt, r)] = formule.format(r=r, FL=R.F_LIGNE, FE=R.F_EQUIPE)
        ws.row_dimensions[r].height = 15

    cause = {c[0]: c for c in D.CAUSES}
    for i, a in enumerate(arrets):
        r = FIRST + i
        if r > LAST:
            break
        ws["A%d" % r] = a["date"]
        ws["A%d" % r].number_format = DATE
        ws["B%d" % r] = a["equipe"]
        ws["C%d" % r] = a["ligne"]
        ws["D%d" % r] = a["machine"]
        ws["E%d" % r] = a["debut"]
        ws["F%d" % r] = a["fin"]
        ws["H%d" % r] = a["code"]

    def dv(kind, rng, **kw):
        v = DataValidation(allow_blank=True, showErrorMessage=True, **kw)
        v.type = kind
        ws.add_data_validation(v)
        v.add(rng)

    for col, name in [("B", "LST_EQUIPES"), ("C", "LST_LIGNES"), ("H", "LST_CAUSES")]:
        dv("list", "%s%d:%s%d" % (col, FIRST, col, LAST), formula1=name,
           errorTitle="Valeur hors référentiel",
           error="Choisir une valeur de la liste (onglet PARAMETRES).")
    dv("date", "A%d:A%d" % (FIRST, LAST), operator="between",
       formula1="DATE(2020,1,1)", formula2="DATE(2099,12,31)",
       errorTitle="Date invalide", error="Saisir une date valide.")

    ws.conditional_formatting.add(
        "A%d:Q%d" % (FIRST, LAST),
        FormulaRule(formula=['$G%d>=60' % FIRST], fill=fill(RED_BG), font=f(9, True, RED)))
    ws.conditional_formatting.add(
        "A%d:Q%d" % (FIRST, LAST),
        FormulaRule(formula=['AND($G%d>=30,$G%d<60)' % (FIRST, FIRST)],
                    fill=fill(AMBER_BG), font=f(9, color=AMBER)))
    ws.conditional_formatting.add(
        "K%d:K%d" % (FIRST, LAST),
        FormulaRule(formula=['$K%d="Planifié"' % FIRST], font=f(9, color=GREY)))

    ws.freeze_panes = "E5"
    ws.auto_filter.ref = "A4:Q%d" % LAST
    page(ws, "landscape", title_rows="4:4")
    ws.sheet_view.zoomScale = 90
    return ws
