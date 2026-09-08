# -*- coding: utf-8 -*-
"""Onglet PARAMETRES : identification, objectifs, referentiels, listes."""
from openpyxl.styles import Alignment, Font, Border, Side
from openpyxl.worksheet.datavalidation import DataValidation
from openpyxl.workbook.defined_name import DefinedName
from openpyxl.utils import get_column_letter as gcl
from common import *
import data as D

R0 = 7          # premiere ligne de donnees des referentiels
NL, NEQ, NPIL = 30, 30, 30
NPROD, NCAUSE, NPOSTE, NLST = 30, 40, 20, 20


def build(wb):
    ws = wb.create_sheet("PARAMETRES")
    ws.sheet_view.showGridLines = False
    ws.sheet_properties.tabColor = "7F8C9A"

    title_band(ws, 1, 7, "PARAMÈTRES ET RÉFÉRENTIELS DU SITE",
               "Source unique de vérité du classeur : objectifs, listes de choix et données de référence. "
               "Toute modification se répercute automatiquement sur les onglets de saisie et le cockpit.")

    # ------------------------------------------------ 1. Identification
    band(ws, 4, 1, 7, "1.  IDENTIFICATION DE L'ENTITÉ PILOTÉE")
    ident = [
        ("Site / usine", "Site de démonstration", "Texte libre — figure dans l'en-tête du cockpit"),
        ("Atelier / UAP", "UAP Assemblage-Usinage", "Périmètre couvert par ce classeur"),
        ("Responsable de production", "À compléter", ""),
        ("Superviseur référent", "À compléter", ""),
        ("Nombre d'équipes par jour", 2, "Sert au dimensionnement des ressources"),
        ("Temps d'ouverture théorique par poste (min)", 480, "Valeur par défaut proposée à la saisie"),
        ("Jours ouvrés de référence par mois", 21, "Base de calcul des ratios mensuels"),
        ("Devise de valorisation", "EUR", "Gains du plan d'actions"),
    ]
    for i, (lab, val, com) in enumerate(ident):
        r = 5 + i
        label(ws, "A%d" % r, lab, bold=False)
        ws.merge_cells("A%d:B%d" % (r, r))
        c = ws.cell(row=r, column=3, value=val)
        c.font = f(10, True, "0000FF")
        c.fill = fill(INPUT_BG)
        c.border = BOX
        c.alignment = Alignment(horizontal="left", vertical="center", indent=1)
        ws.merge_cells("C%d:E%d" % (r, r))
        for cc in range(4, 6):
            ws.cell(row=r, column=cc).fill = fill(INPUT_BG)
            ws.cell(row=r, column=cc).border = BOX
        note(ws, "F%d" % r, com, "F%d:G%d" % (r, r))
        ws.row_dimensions[r].height = 17

    # ------------------------------------------------ 2. Objectifs
    band(ws, 14, 1, 7, "2.  OBJECTIFS ET SEUILS D'ALERTE  —  contrat de performance")
    header_row(ws, 15, ["Code", "Indicateur", "Unité", "Cible", "Seuil d'alerte",
                        "Sens", "Origine de l'objectif"], height=30)
    for i, (code, lib, unite, cible, alerte, sens, src) in enumerate(D.OBJECTIFS):
        r = 16 + i
        vals = [code, lib, unite, cible, alerte, sens, src]
        for j, v in enumerate(vals):
            c = ws.cell(row=r, column=1 + j, value=v)
            c.border = BOX
            c.font = f(9)
            c.alignment = Alignment(horizontal="left", vertical="center", indent=1)
        for col in (4, 5):
            c = ws.cell(row=r, column=col)
            c.font = f(9, True, "0000FF")
            c.fill = fill(INPUT_BG)
            c.alignment = Alignment(horizontal="center", vertical="center")
            c.number_format = PCT if unite == "%" else NUM
        ws.cell(row=r, column=1).font = f(9, True, STEEL)
        ws.cell(row=r, column=1).alignment = Alignment(horizontal="center", vertical="center")
        ws.cell(row=r, column=3).alignment = Alignment(horizontal="center", vertical="center")
        ws.cell(row=r, column=6).alignment = Alignment(horizontal="center", vertical="center")
        ws.cell(row=r, column=7).font = f(8.5, False, GREY)
        ws.row_dimensions[r].height = 16
    zebra(ws, 16, 15 + len(D.OBJECTIFS), 1, 7)
    note(ws, "A%d" % (17 + len(D.OBJECTIFS)),
         "Sens de lecture :  1 = la performance s'améliore quand la valeur augmente  •  "
         "-1 = la performance s'améliore quand la valeur diminue. "
         "Le cockpit compare automatiquement la valeur du mois à la cible (vert), au seuil d'alerte (orange) puis au-delà (rouge).",
         "A%d:G%d" % (17 + len(D.OBJECTIFS), 18 + len(D.OBJECTIFS)))

    # ------------------------------------------------ Referentiels (zone droite)
    band(ws, 4, 9, 38, "3.  RÉFÉRENTIELS DE SAISIE  —  alimentent les listes déroulantes de tous les onglets")

    def block(c1, c2, titre, headers, rows, values=None, fmts=None):
        ws.merge_cells(start_row=5, start_column=c1, end_row=5, end_column=c2)
        t = ws.cell(row=5, column=c1, value=titre)
        t.font = f(9.5, True, WHITE)
        t.alignment = Alignment(horizontal="center", vertical="center")
        for c in range(c1, c2 + 1):
            ws.cell(row=5, column=c).fill = fill(SLATE)
        header_row(ws, 6, headers, start_col=c1, bg=STEEL, height=28)
        for i in range(rows):
            r = R0 + i
            for j in range(c2 - c1 + 1):
                c = ws.cell(row=r, column=c1 + j)
                c.border = BOX
                c.font = f(9, color="0000FF")
                c.fill = fill(REF_BG)
                c.alignment = Alignment(horizontal="left", vertical="center", indent=1)
                if fmts and fmts[j]:
                    c.number_format = fmts[j]
                    c.alignment = Alignment(horizontal="center", vertical="center")
            ws.row_dimensions[r].height = 15
        if values:
            for i, row in enumerate(values):
                for j, v in enumerate(row):
                    if v is not None:
                        ws.cell(row=R0 + i, column=c1 + j, value=v)

    block(9, 11, "Organisation", ["Lignes / ressources", "Équipes (postes)", "Pilotes et animateurs"], NL,
          [[D.LIGNES[i] if i < len(D.LIGNES) else None,
            D.EQUIPES[i] if i < len(D.EQUIPES) else None,
            D.PILOTES[i] if i < len(D.PILOTES) else None] for i in range(NL)])
    block(13, 16, "Produits et cadences nominales",
          ["Référence", "Désignation", "Cadence nominale\n(pcs/min)", "PPM cible"], NPROD,
          [[p[0], p[1], p[2], p[3]] for p in D.PRODUITS], fmts=[None, None, "0.0", NUM])
    block(18, 22, "Causes d'arrêt normalisées",
          ["Code", "Libellé", "Famille 5M", "Type", "Rubrique de saisie"], NCAUSE,
          [list(c) for c in D.CAUSES])
    block(24, 26, "Postes de travail (polyvalence)",
          ["Poste de travail", "Criticité\n(1 à 3)", "Effectif autonome\ncible"], NPOSTE,
          [[D.POSTES[i], D.POSTES_CRIT[i], D.POSTES_CIBLE[i]] for i in range(len(D.POSTES))],
          fmts=[None, "0", "0"])
    listes = [
        ("Familles 5M", D.FAM5M), ("Sources d'écart", D.SOURCES),
        ("Nature d'action", D.NATURES), ("Statut d'action", D.STATUTS),
        ("Efficacité", D.EFFICACITE), ("Oui / Non", D.OUINON),
        ("Contrôle terrain", D.OKNOK), ("Axes SQCDP", D.SQCDP),
        ("Filtre lignes", ["Toutes"] + D.LIGNES),
        ("Filtre équipes", ["Toutes"] + D.EQUIPES),
        ("Rubriques de perte", D.RUBRIQUES), ("Mois", D.MOIS),
    ]
    block(28, 39, "Listes de choix normalisées", [l[0] for l in listes], NLST,
          [[listes[j][1][i] if i < len(listes[j][1]) else None for j in range(len(listes))]
           for i in range(NLST)])

    widths(ws, {"A": 22, "B": 24, "C": 12, "D": 12, "E": 13, "F": 8, "G": 46,
                "I": 26, "J": 22, "K": 20, "L": 2.5, "M": 12, "N": 26, "O": 12, "P": 11,
                "Q": 2.5, "R": 10, "S": 36, "T": 14, "U": 11, "V": 24, "W": 2.5,
                "X": 24, "Y": 10, "Z": 12, "AA": 2.5})
    for i in range(28, 40):
        ws.column_dimensions[gcl(i)].width = 19
    ws.freeze_panes = "A7"

    # ------------------------------------------------ Noms definis
    def dyn(name, col, first, count):
        ref = "OFFSET(PARAMETRES!${c}${f},0,0,MAX(1,COUNTA(PARAMETRES!${c}${f}:${c}${l})),1)".format(
            c=col, f=first, l=first + count - 1)
        wb.defined_names.add(DefinedName(name, attr_text=ref))

    dyn("LST_LIGNES", "I", R0, NL)
    dyn("LST_EQUIPES", "J", R0, NEQ)
    dyn("LST_PILOTES", "K", R0, NPIL)
    dyn("LST_PRODUITS", "M", R0, NPROD)
    dyn("LST_CAUSES", "R", R0, NCAUSE)
    dyn("LST_POSTES", "X", R0, NPOSTE)
    for name, col in [("LST_FAM5M", "AB"), ("LST_SOURCES", "AC"), ("LST_NATURES", "AD"),
                      ("LST_STATUTS", "AE"), ("LST_EFFICACITE", "AF"), ("LST_OUINON", "AG"),
                      ("LST_CONTROLE", "AH"), ("LST_SQCDP", "AI"), ("LST_F_LIGNES", "AJ"),
                      ("LST_F_EQUIPES", "AK"), ("LST_RUBRIQUES", "AL"), ("LST_MOIS", "AM")]:
        dyn(name, col, R0, NLST)

    dv = DataValidation(type="whole", operator="between", formula1="1", formula2="3",
                        allow_blank=True, showErrorMessage=True,
                        errorTitle="Criticité", error="Saisir un entier de 1 à 3.")
    ws.add_data_validation(dv)
    dv.add("Y%d:Y%d" % (R0, R0 + NPOSTE - 1))

    page(ws, "landscape", title_rows="1:2")
    return ws
