# -*- coding: utf-8 -*-
"""Onglet COCKPIT : tableau de bord de la revue de performance."""
from openpyxl.styles import Alignment, Border, Side
from openpyxl.worksheet.datavalidation import DataValidation
from openpyxl.formatting.rule import FormulaRule
from openpyxl.chart import LineChart, BarChart, DoughnutChart, Reference
from openpyxl.chart.series import SeriesLabel
from openpyxl.chart.shapes import GraphicalProperties
from openpyxl.chart.marker import Marker
from openpyxl.drawing.line import LineProperties
from openpyxl.chart.label import DataLabelList
from common import *
import s_calc as C

NC = 18
CIBLE = 'INDEX(PARAMETRES!$D$16:$D$29,MATCH("{c}",PARAMETRES!$A$16:$A$29,0))'
ALERTE = 'INDEX(PARAMETRES!$E$16:$E$29,MATCH("{c}",PARAMETRES!$A$16:$A$29,0))'

# (code, titre, formule, format cellule, format TEXT, suffixe, colonne M-1 dans CALC)
CARTES = [
    ("TRS",      "TRS DU MOIS\n(rendement synthétique)",
     '=IFERROR(CALC!$AC$20/CALC!$AC$18,"")', PCT, '0.0%', "", "L"),
    ("DISPO",    "DISPONIBILITÉ",
     '=IFERROR(CALC!$AC$19/CALC!$AC$18,"")', PCT, '0.0%', "", "M"),
    ("PERF",     "PERFORMANCE\n(tenue de cadence)",
     '=IFERROR((CALC!$AC$19-CALC!$AC$16)/CALC!$AC$19,"")', PCT, '0.0%', "", "N"),
    ("QUAL",     "QUALITÉ AU PREMIER PASSAGE",
     '=IFERROR(CALC!$AC$23/CALC!$AC$22,"")', PCT, '0.0%', "", "O"),
    ("TRG",      "TRG\n(sur temps d'ouverture)",
     '=IFERROR(CALC!$AC$20/CALC!$AC$8,"")', PCT, '0.0%', "", "P"),
    ("SERVICE",  "TAUX DE SERVICE",
     '=IFERROR((CALC!$AC$23+CALC!$AC$24)/CALC!$AC$25,"")', PCT, '0.0%', "", "Q"),
    ("PPM",      "NON-CONFORMITÉS INTERNES",
     '=IFERROR((CALC!$AC$22-CALC!$AC$23)/CALC!$AC$22*1000000,"")', PPM, '#,##0', " ppm", None),
    ("PRESENCE", "TAUX DE PRÉSENCE",
     '=IFERROR(CALC!$AC$29/CALC!$AC$28,"")', PCT, '0.0%', "", None),
    ("ACCIDENT", "ACCIDENTS AVEC ARRÊT",
     '=CALC!$AC$26', NUM, '0', "", None),
    ("MTBF",     "MTBF\n(temps entre pannes)",
     '=IFERROR(CALC!$AC$19/CALC!$AC$27,"")', MIN, '#,##0', " min", None),
    ("MTTR",     "MTTR\n(temps de réparation)",
     '=IFERROR(CALC!$AC$10/CALC!$AC$27,"")', MIN, '#,##0', " min", None),
    ("RETARD",   "ACTIONS EN RETARD",
     '=PLAN_ACTIONS!$D$6', NUM, '0', "", None),
]

TOP_HDR = ["Rang", "Code", "Cause d'arrêt", "Minutes perdues", "Occurrences",
           "Durée moyenne", "Part des arrêts subis"]


def build(wb):
    ws = wb.create_sheet("COCKPIT")
    ws.sheet_view.showGridLines = False
    ws.sheet_properties.tabColor = "0F2A44"

    title_band(ws, 1, NC, "COCKPIT DE PILOTAGE DE LA PRODUCTION",
               '=PARAMETRES!$C$5&"  —  "&PARAMETRES!$C$6&"   |   Revue quotidienne de performance   |   '
               'Édition du "&TEXT(TODAY(),"dd/mm/yyyy")')
    ws["A2"].font = f(9, False, "4A5A6A", italic=True)

    # ---------------------------------------------------------- filtres
    for col in range(1, NC + 1):
        cc = ws.cell(row=4, column=col)
        cc.fill = fill(LIGHT)
        cc.border = Border(bottom=Side(style="thin", color=STEEL))
    filtres = [("A4", "Année :", "B4", 2026, "0", None),
               ("C4", "Mois :", "D4", 1, "0", None),
               ("G4", "Ligne :", "H4", "Toutes", None, "LST_F_LIGNES"),
               ("J4", "Équipe :", "K4", "Toutes", None, "LST_F_EQUIPES")]
    for lab_ref, lab, val_ref, val, fmt, lst in filtres:
        c = ws[lab_ref]
        c.value = lab
        c.font = f(9.5, True, STEEL)
        c.alignment = Alignment(horizontal="right", vertical="center")
        v = ws[val_ref]
        v.value = val
        v.font = f(10, True, "0000FF")
        v.fill = fill(INPUT_BG)
        v.border = BOX
        v.alignment = Alignment(horizontal="center", vertical="center")
        if fmt:
            v.number_format = fmt
        if lst:
            dv = DataValidation(type="list", formula1=lst, allow_blank=False)
            ws.add_data_validation(dv)
            dv.add(val_ref)
    ws.merge_cells("H4:I4")
    ws.merge_cells("K4:L4")
    for ref in ("I4", "L4"):
        ws[ref].fill = fill(INPUT_BG)
        ws[ref].border = BOX
    ws.merge_cells("E4:F4")
    e = ws["E4"]
    e.value = '=INDEX(PARAMETRES!$AM$7:$AM$18,$D$4)&" "&TEXT($B$4,"0000")'
    e.font = f(10, True, NAVY)
    e.alignment = Alignment(horizontal="left", vertical="center", indent=1)
    ws.merge_cells("N4:R4")
    n = ws["N4"]
    n.value = ('="Jours de production saisis : "&CALC!$B$9&"   |   Lignes suivies : "'
               '&COUNTIF(CALC!$D$%d:$D$%d,">0")' % (C.LIG0, C.LIGN))
    n.font = f(9, False, GREY, italic=True)
    n.alignment = Alignment(horizontal="right", vertical="center")
    ws.row_dimensions[4].height = 22

    # ---------------------------------------------------------- cartes
    for idx, (code, titre, formule, fmt, tfmt, suffixe, m1) in enumerate(CARTES):
        row = 6 if idx < 6 else 12
        c1 = 1 + (idx % 6) * 3
        c2 = c1 + 2
        hr = 6 + idx                      # ligne du bloc technique
        suf = ('&"%s"' % suffixe) if suffixe else ""
        cible_txt = '="Cible : "&TEXT($V$%d,"%s")%s' % (hr, tfmt, suf)
        if m1:
            ref = "CALC!$%s$%d" % (m1, C.MOIN - 1)
            sous2 = ('=IF(%s="","Mois précédent : non disponible","M-1 : "&TEXT(%s,"0.0%%")&'
                     '"    "&TEXT(($U$%d-%s)*100,"+0.0 pt;-0.0 pt;0.0 pt"))'
                     % (ref, ref, hr, ref))
        else:
            sous2 = '="Seuil d\'alerte : "&TEXT($W$%d,"%s")%s' % (hr, tfmt, suf)
        kpi_tile(ws, row, c1, c2, titre, formule, fmt, cible_txt, sous2)

        # bloc technique (colonnes T a Y, masquees)
        ws["T%d" % hr] = code
        ws["U%d" % hr] = "=%s%d" % (chr(64 + c1) if c1 < 27 else "A", row + 1)
        ws["V%d" % hr] = "=" + CIBLE.format(c=code)
        ws["W%d" % hr] = "=" + ALERTE.format(c=code)
        ws["X%d" % hr] = ('=INDEX(PARAMETRES!$F$16:$F$29,MATCH("%s",PARAMETRES!$A$16:$A$29,0))'
                          % code)
        ws["Y%d" % hr] = ('=IF($U{r}="","",IF($X{r}=1,IF($U{r}>=$V{r},1,IF($U{r}<$W{r},3,2)),'
                          'IF($U{r}<=$V{r},1,IF($U{r}>$W{r},3,2))))').format(r=hr)
        for col in "TUVWXY":
            ws["%s%d" % (col, hr)].font = f(8, color=GREY)

        val_rng = "%s%d:%s%d" % (chr(64 + c1), row + 1, chr(64 + c2), row + 2)
        for statut, bg, fg in ((1, GREEN_BG, GREEN), (2, AMBER_BG, AMBER), (3, RED_BG, RED)):
            ws.conditional_formatting.add(
                val_rng, FormulaRule(formula=['$Y$%d=%d' % (hr, statut)],
                                     fill=fill(bg), font=f(20, True, fg)))
        ws.conditional_formatting.add(
            "%s%d:%s%d" % (chr(64 + c1), row, chr(64 + c2), row),
            FormulaRule(formula=['$Y$%d=3' % hr], fill=fill(RED), font=f(8.5, True, WHITE)))

    label(ws, "T5", "Bloc technique — pilotage des couleurs (ne pas modifier)",
          bold=True, size=8, color=GREY)
    for col in "TUVWXY":
        ws.column_dimensions[col].hidden = True

    # ---------------------------------------------------------- graphiques
    band(ws, 18, 1, NC, "ANALYSE DE LA PERFORMANCE DU MOIS")
    _charts(ws)

    # ---------------------------------------------------------- causes et décisions
    band(ws, 78, 1, NC, "CAUSES D'ARRÊT SUBI DU MOIS  ·  DÉCISIONS DE LA REVUE")
    _chart_causes(ws)

    dspans = [(10, 14), (15, 17), (18, 18)]
    for j, titre in enumerate(["Décision prise", "Pilote", "Échéance"]):
        c1, c2 = dspans[j]
        ws.merge_cells(start_row=79, start_column=c1, end_row=79, end_column=c2)
        c = ws.cell(row=79, column=c1, value=titre)
        c.font = f(8.5, True, WHITE)
        c.alignment = Alignment(horizontal="center", vertical="center")
        for k in range(c1, c2 + 1):
            ws.cell(row=79, column=k).fill = fill(SLATE)
            ws.cell(row=79, column=k).border = BOX
    ws.row_dimensions[79].height = 22
    for i in range(5):
        r = 80 + i
        for c1, c2 in dspans:
            ws.merge_cells(start_row=r, start_column=c1, end_row=r, end_column=c2)
            c = ws.cell(row=r, column=c1)
            c.font = f(9, color="0000FF")
            c.alignment = Alignment(horizontal="left", vertical="center", indent=1)
            for k in range(c1, c2 + 1):
                ws.cell(row=r, column=k).fill = fill(INPUT_BG)
                ws.cell(row=r, column=k).border = BOX
        ws.cell(row=r, column=18).number_format = DATE
        ws.row_dimensions[r].height = 22
    dv = DataValidation(type="list", formula1="LST_PILOTES", allow_blank=True)
    ws.add_data_validation(dv)
    dv.add("O80:O84")
    note(ws, "J86",
         "Le graphique de gauche classe les causes d'arrêt subi du mois par minutes perdues. "
         "Chaque cause récurrente doit se retrouver dans le plan d'actions avec un pilote et une date.",
         "J86:R88")

    note(ws, "A92",
         "Lecture des cartes :  vert = cible atteinte  •  orange = entre la cible et le seuil d'alerte  •  "
         "rouge = au-delà du seuil d'alerte.  Cibles et seuils sont définis dans l'onglet PARAMETRES. "
         "Règle d'animation : toute carte rouge donne lieu, le jour même, à une action ouverte dans le PLAN_ACTIONS "
         "avec un pilote nommé et une date cible. Le graphique de suivi journalier affiche la moyenne mobile 7 jours "
         "et les limites de contrôle (moyenne ± 2 écarts-types) : un point hors limite signale une cause spéciale à investiguer, "
         "à distinguer d'une variation courante du procédé.",
         "A92:R94")

    for col in range(1, NC + 1):
        ws.column_dimensions[chr(64 + col) if col < 27 else "A"].width = 9.6
    ws.freeze_panes = "A5"
    page(ws, "landscape", area="A1:R94")
    ws.sheet_view.zoomScale = 80
    return ws


def _chart_causes(ws):
    calc = ws.parent["CALC"]
    b = BarChart()
    b.type = "bar"
    b.title = "Top 5 des causes d'arrêt subi (minutes perdues)"
    b.height, b.width = 8.4, 14.6
    b.add_data(Reference(calc, min_col=4, min_row=C.TOP0, max_row=C.TOPN), titles_from_data=False)
    b.set_categories(Reference(calc, min_col=3, min_row=C.TOP0, max_row=C.TOPN))
    str_categories(b, "'CALC'!$C$%d:$C$%d" % (C.TOP0, C.TOPN))
    b.series[0].tx = SeriesLabel(v="Minutes perdues")
    b.series[0].graphicalProperties = GraphicalProperties(solidFill="C0392B")
    b.gapWidth = 40
    b.legend = None
    b.dLbls = DataLabelList()
    b.dLbls.showVal = True
    b.x_axis.majorGridlines = None
    ws.add_chart(b, "A79")


def _serie(ch, i, color, width=20000, dash=None, marker=False, smooth=False):
    s = ch.series[i]
    s.graphicalProperties = GraphicalProperties()
    lp = LineProperties(solidFill=color, w=width)
    if dash:
        lp.prstDash = dash
    s.graphicalProperties.line = lp
    s.smooth = smooth
    if not marker:
        s.marker = Marker(symbol="none")
    return s


def _charts(ws):
    # 1. suivi journalier du TRS
    ch = LineChart()
    ch.title = "Suivi journalier du TRS — moyenne mobile 7 jours et limites de contrôle"
    ch.height, ch.width = 10.2, 16.6
    ch.style = 2
    for col in ("R", "S", "N", "P", "Q"):
        idx = {"R": 18, "S": 19, "N": 14, "P": 16, "Q": 17}[col]
        ch.add_data(Reference(ws.parent["CALC"], min_col=idx, min_row=C.JOUR0, max_row=C.JOURN),
                    titles_from_data=False)
    ch.set_categories(Reference(ws.parent["CALC"], min_col=1, min_row=C.JOUR0, max_row=C.JOURN))
    for i, (lib, color, dash, w) in enumerate([
            ("TRS journalier", "1F6FB2", None, 22000),
            ("Moyenne mobile 7 jours", "B5651D", None, 22000),
            ("Cible", "1E7B34", "dash", 15000),
            ("Limite de contrôle supérieure", "9AA7B2", "sysDot", 10000),
            ("Limite de contrôle inférieure", "9AA7B2", "sysDot", 10000)]):
        ch.series[i].tx = SeriesLabel(v=lib)
        _serie(ch, i, color, w, dash)
    ch.y_axis.numFmt = "0%"
    ch.y_axis.title = "TRS"
    ch.x_axis.title = "Jour du mois"
    ch.y_axis.majorGridlines = None
    ws.add_chart(ch, "A19")

    # 2. Pareto des pertes de temps
    bar = BarChart()
    bar.type = "col"
    bar.title = "Pareto des pertes de temps subies (minutes)"
    bar.height, bar.width = 10.2, 16.6
    bar.add_data(Reference(ws.parent["CALC"], min_col=7, min_row=C.PAR0, max_row=C.PARN),
                 titles_from_data=False)
    bar.set_categories(Reference(ws.parent["CALC"], min_col=6, min_row=C.PAR0, max_row=C.PARN))
    bar.series[0].tx = SeriesLabel(v="Minutes perdues")
    bar.series[0].graphicalProperties = GraphicalProperties(solidFill="1F6FB2")
    bar.gapWidth = 45
    bar.y_axis.title = "Minutes"
    bar.y_axis.majorGridlines = None
    cum = LineChart()
    cum.add_data(Reference(ws.parent["CALC"], min_col=9, min_row=C.PAR0, max_row=C.PARN),
                 titles_from_data=False)
    cum.series[0].tx = SeriesLabel(v="Part cumulée")
    _serie(cum, 0, "B01919", 22000, marker=True)
    cum.series[0].marker = Marker(symbol="circle", size=5)
    cum.y_axis.axId = 300
    cum.y_axis.numFmt = "0%"
    cum.y_axis.title = "Cumul"
    cum.y_axis.scaling.max = 1
    cum.y_axis.majorGridlines = None
    cum.y_axis.crosses = "max"
    bar.y_axis.crosses = "autoZero"
    str_categories(bar, "'CALC'!$F$%d:$F$%d" % (C.PAR0, C.PARN))
    str_categories(cum, "'CALC'!$F$%d:$F$%d" % (C.PAR0, C.PARN))
    bar += cum
    ws.add_chart(bar, "J19")

    # 3. cascade des pertes de temps
    wf = BarChart()
    wf.type = "col"
    wf.grouping = "stacked"
    wf.overlap = 100
    wf.title = "Cascade des pertes : du temps d'ouverture au temps utile (minutes)"
    wf.height, wf.width = 10.2, 16.6
    wf.add_data(Reference(ws.parent["CALC"], min_col=2, min_row=C.CAS0, max_row=C.CASN),
                titles_from_data=False)
    wf.add_data(Reference(ws.parent["CALC"], min_col=3, min_row=C.CAS0, max_row=C.CASN),
                titles_from_data=False)
    wf.set_categories(Reference(ws.parent["CALC"], min_col=1, min_row=C.CAS0, max_row=C.CASN))
    wf.series[0].tx = SeriesLabel(v="Socle")
    wf.series[1].tx = SeriesLabel(v="Minutes")
    wf.series[0].graphicalProperties = GraphicalProperties(noFill=True)
    wf.series[0].graphicalProperties.line = LineProperties(noFill=True)
    wf.series[1].graphicalProperties = GraphicalProperties(solidFill="35566F")
    wf.gapWidth = 40
    wf.y_axis.title = "Minutes"
    wf.y_axis.majorGridlines = None
    wf.legend = None
    str_categories(wf, "'CALC'!$A$%d:$A$%d" % (C.CAS0, C.CASN))
    ws.add_chart(wf, "A39")

    # 4. historique 13 mois
    hi = LineChart()
    hi.title = "Tendance du TRS sur 13 mois glissants"
    hi.height, hi.width = 10.2, 16.6
    hi.add_data(Reference(ws.parent["CALC"], min_col=18, min_row=C.MOI0, max_row=C.MOIN),
                titles_from_data=False)
    hi.add_data(Reference(ws.parent["CALC"], min_col=19, min_row=C.MOI0, max_row=C.MOIN),
                titles_from_data=False)
    hi.set_categories(Reference(ws.parent["CALC"], min_col=1, min_row=C.MOI0, max_row=C.MOIN))
    hi.series[0].tx = SeriesLabel(v="TRS mensuel")
    hi.series[1].tx = SeriesLabel(v="Cible")
    _serie(hi, 0, "1F6FB2", 24000, marker=True)
    hi.series[0].marker = Marker(symbol="circle", size=5)
    _serie(hi, 1, "1E7B34", 15000, "dash")
    str_categories(hi, "'CALC'!$A$%d:$A$%d" % (C.MOI0, C.MOIN))
    hi.y_axis.numFmt = "0%"
    hi.y_axis.majorGridlines = None
    ws.add_chart(hi, "J39")

    # 5 et 6. TRS par ligne et par equipe
    for anchor, (r0, rn), titre, color in (("A58", (C.LIG0, C.LIGN), "TRS par ligne de production", "1F6FB2"),
                                           ("G58", (C.EQU0, C.EQUN), "TRS par équipe", "35566F")):
        b = BarChart()
        b.type = "bar"
        b.title = titre + " — mois sélectionné"
        b.height, b.width = 9.0, 10.6
        b.add_data(Reference(ws.parent["CALC"], min_col=5, min_row=r0, max_row=rn),
                   titles_from_data=False)
        b.set_categories(Reference(ws.parent["CALC"], min_col=1, min_row=r0, max_row=rn))
        b.series[0].tx = SeriesLabel(v="TRS")
        b.series[0].graphicalProperties = GraphicalProperties(solidFill=color)
        b.x_axis.numFmt = "0%"
        b.gapWidth = 55
        b.legend = None
        b.dLbls = DataLabelList()
        b.dLbls.showVal = True
        b.dLbls.numFmt = "0.0%"
        str_categories(b, "'CALC'!$A$%d:$A$%d" % (r0, rn))
        ws.add_chart(b, anchor)

    # 7. avancement du plan d'actions
    dn = DoughnutChart(holeSize=52)
    dn.title = "Avancement du plan d'actions"
    dn.height, dn.width = 9.0, 10.6
    dn.add_data(Reference(ws.parent["CALC"], min_col=2, min_row=C.ACT0, max_row=C.ACTN),
                titles_from_data=False)
    dn.set_categories(Reference(ws.parent["CALC"], min_col=1, min_row=C.ACT0, max_row=C.ACTN))
    dn.series[0].tx = SeriesLabel(v="Actions")
    dn.dLbls = DataLabelList()
    dn.dLbls.showVal = True
    str_categories(dn, "'CALC'!$A$%d:$A$%d" % (C.ACT0, C.ACTN))
    ws.add_chart(dn, "M58")
