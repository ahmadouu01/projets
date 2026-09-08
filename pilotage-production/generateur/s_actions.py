# -*- coding: utf-8 -*-
"""Onglet PLAN_ACTIONS : resolution de probleme structuree (QRQC / 8D)."""
import datetime as dt
from openpyxl.styles import Alignment, Border, Side
from openpyxl.worksheet.datavalidation import DataValidation
from openpyxl.formatting.rule import FormulaRule, ColorScaleRule
from common import *
import data as D

FIRST, LAST = 9, 128

COLS = [
    ("A", "N°", 8, "ca", None),
    ("B", "Date de\ndétection", 11, "in", DATE),
    ("C", "Source", 18, "in", None),
    ("D", "Ligne / zone", 20, "in", None),
    ("E", "Problème constaté (fait mesuré)", 44, "in", None),
    ("F", "G", 5, "in", "0"),
    ("G", "F", 5, "in", "0"),
    ("H", "D", 5, "in", "0"),
    ("I", "Criticité\n(IPR)", 9, "ca", "0"),
    ("J", "Famille 5M", 14, "in", None),
    ("K", "Cause racine validée (5 pourquoi)", 44, "in", None),
    ("L", "Action décidée", 44, "in", None),
    ("M", "Nature", 15, "in", None),
    ("N", "Pilote", 14, "in", None),
    ("O", "Date cible", 11, "in", DATE),
    ("P", "Avanc.", 8, "in", PCT0),
    ("Q", "Statut", 13, "in", None),
    ("R", "Date de\nclôture", 11, "in", DATE),
    ("S", "Retard\n(j)", 8, "ca", NUM),
    ("T", "Délai de\ntraitement (j)", 10, "ca", NUM),
    ("U", "Efficacité\nvérifiée", 11, "in", None),
    ("V", "Date de\nvérification", 11, "in", DATE),
    ("W", "Standard\nmis à jour", 10, "in", None),
    ("X", "Gain estimé\n(€/an)", 12, "in", EUR),
    ("Y", "Commentaire de suivi", 34, "in", None),
]

FORMULES = {
    "A": '=IF($E{r}="","","A-"&TEXT(ROW()-{off},"000"))',
    "I": '=IF(OR($F{r}="",$G{r}="",$H{r}=""),"",$F{r}*$G{r}*$H{r})',
    "S": ('=IF($O{r}="","",IF(OR($Q{r}="Terminée",$Q{r}="Abandonnée"),'
          'IF($R{r}="","",MAX(0,$R{r}-$O{r})),MAX(0,TODAY()-$O{r})))'),
    "T": '=IF(OR($R{r}="",$B{r}=""),"",$R{r}-$B{r})',
}

TUILES = [
    ("Actions ouvertes", 1, 3,
     '=COUNTIF($Q${f}:$Q${l},"À lancer")+COUNTIF($Q${f}:$Q${l},"En cours")+COUNTIF($Q${f}:$Q${l},"En attente")', NUM),
    ("Dont en retard", 4, 6,
     '=SUMPRODUCT(($S${f}:$S${l}<>"")*($S${f}:$S${l}>0)*($Q${f}:$Q${l}<>"Terminée")*($Q${f}:$Q${l}<>"Abandonnée"))', NUM),
    ("Respect des délais", 7, 9,
     '=IFERROR(COUNTIFS($Q${f}:$Q${l},"Terminée",$S${f}:$S${l},0)/COUNTIF($Q${f}:$Q${l},"Terminée"),"")', PCT),
    ("Délai moyen de traitement", 10, 12, '=IFERROR(AVERAGE($T${f}:$T${l}),"")', '#,##0.0" j"'),
    ("Criticité moyenne (IPR)", 13, 15, '=IFERROR(AVERAGE($I${f}:$I${l}),"")', "0.0"),
    ("Actions critiques (IPR ≥ 30)", 16, 18, '=COUNTIF($I${f}:$I${l},">=30")', NUM),
    ("Efficacité confirmée", 19, 21,
     '=IFERROR(COUNTIF($U${f}:$U${l},"Oui")/(COUNTIF($U${f}:$U${l},"Oui")+COUNTIF($U${f}:$U${l},"Non")),"")', PCT),
    ("Gains sécurisés (€/an)", 22, 25, '=SUMIF($Q${f}:$Q${l},"Terminée",$X${f}:$X${l})', EUR),
]


def build(wb):
    ws = wb.create_sheet("PLAN_ACTIONS")
    ws.sheet_view.showGridLines = False
    ws.sheet_properties.tabColor = "B5651D"
    n = len(COLS)

    title_band(ws, 1, n, "PLAN D'ACTIONS  —  résolution de problème structurée",
               "Règle d'animation : aucun problème sans fait mesuré, sans cause racine validée, sans pilote nommé et sans date cible. "
               "Trois chantiers de fond ouverts au maximum. Criticité IPR = Gravité × Fréquence × Détection (échelle 1 à 5).")

    band(ws, 4, 1, n, "SYNTHÈSE DU PLAN D'ACTIONS  —  calculée automatiquement")
    for lib, c1, c2, formule, fmt in TUILES:
        ws.merge_cells(start_row=5, start_column=c1, end_row=5, end_column=c2)
        c = ws.cell(row=5, column=c1, value=lib)
        c.font = f(8.5, True, WHITE)
        c.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        for k in range(c1, c2 + 1):
            ws.cell(row=5, column=k).fill = fill(SLATE)
            ws.cell(row=5, column=k).border = BOX
        ws.merge_cells(start_row=6, start_column=c1, end_row=6, end_column=c2)
        v = ws.cell(row=6, column=c1, value=formule.format(f=FIRST, l=LAST))
        v.font = f(14, True, NAVY)
        v.number_format = fmt
        v.alignment = Alignment(horizontal="center", vertical="center")
        for k in range(c1, c2 + 1):
            ws.cell(row=6, column=k).fill = fill(WHITE)
            ws.cell(row=6, column=k).border = BOX
    ws.row_dimensions[5].height = 26
    ws.row_dimensions[6].height = 26

    header_row(ws, 8, [c[1] for c in COLS], height=40)
    ws["F8"].value = "G"
    for lt, tip in (("F", "Gravité (1 à 5)"), ("G", "Fréquence (1 à 5)"), ("H", "Détection (1 à 5)")):
        ws["%s8" % lt].alignment = Alignment(horizontal="center", vertical="center")

    for i, (lt, _, w, kind, fmt) in enumerate(COLS):
        ws.column_dimensions[lt].width = w
        if kind != "in":
            ws.cell(row=8, column=i + 1).fill = fill(NAVY)
        rng = "%s%d:%s%d" % (lt, FIRST, lt, LAST)
        if kind == "in":
            apply_range(ws, rng, font=f(9, color="0000FF"), fillc=INPUT_BG, numfmt=fmt,
                        align="center" if fmt else "left", wrap=(w > 30))
        else:
            apply_range(ws, rng, font=f(9), fillc=CALC_BG, numfmt=fmt, align="center")

    for r in range(FIRST, LAST + 1):
        for lt, formule in FORMULES.items():
            ws["%s%d" % (lt, r)] = formule.format(r=r, off=FIRST - 1)
        ws.row_dimensions[r].height = 30

    for i, a in enumerate(D.ACTIONS_DEMO):
        r = FIRST + i
        vals = {"B": dt.date.fromisoformat(a["d"]), "C": a["src"], "D": a["zone"],
                "E": a["pb"], "F": a["g"], "G": a["fr"], "H": a["de"], "J": a["fam"],
                "K": a["cause"], "L": a["act"], "M": a["nat"], "N": a["pil"],
                "O": dt.date.fromisoformat(a["cible"]), "P": a["av"], "Q": a["st"],
                "R": dt.date.fromisoformat(a["fait"]) if a["fait"] else None,
                "U": a["eff"],
                "V": dt.date.fromisoformat(a["verif"]) if a["verif"] else None,
                "W": a["std"], "X": a["gain"], "Y": a["com"]}
        for lt, v in vals.items():
            if v is not None:
                ws["%s%d" % (lt, r)] = v

    def dv(kind, rng, **kw):
        v = DataValidation(allow_blank=True, showErrorMessage=True, **kw)
        v.type = kind
        ws.add_data_validation(v)
        v.add(rng)

    for col, name in [("C", "LST_SOURCES"), ("D", "LST_LIGNES"), ("J", "LST_FAM5M"),
                      ("M", "LST_NATURES"), ("N", "LST_PILOTES"), ("Q", "LST_STATUTS"),
                      ("U", "LST_EFFICACITE"), ("W", "LST_OUINON")]:
        dv("list", "%s%d:%s%d" % (col, FIRST, col, LAST), formula1=name,
           errorTitle="Valeur hors référentiel",
           error="Choisir une valeur de la liste (onglet PARAMETRES).")
    for col in "FGH":
        dv("whole", "%s%d:%s%d" % (col, FIRST, col, LAST), operator="between",
           formula1="1", formula2="5", errorTitle="Cotation",
           error="Coter de 1 (faible) à 5 (fort).")
    dv("decimal", "P%d:P%d" % (FIRST, LAST), operator="between", formula1="0", formula2="1",
       errorTitle="Avancement", error="Saisir un pourcentage entre 0 % et 100 %.")

    ws.conditional_formatting.add(
        "I%d:I%d" % (FIRST, LAST),
        ColorScaleRule(start_type="num", start_value=1, start_color="D6ECD8",
                       mid_type="num", mid_value=30, mid_color="FDE9A9",
                       end_type="num", end_value=80, end_color="F5A6A6"))
    ws.conditional_formatting.add(
        "A%d:Y%d" % (FIRST, LAST),
        FormulaRule(formula=['AND($S%d<>"",$S%d>0,$Q%d<>"Terminée",$Q%d<>"Abandonnée")'
                             % (FIRST, FIRST, FIRST, FIRST)],
                    fill=fill(RED_BG), stopIfTrue=False))
    ws.conditional_formatting.add(
        "Q%d:Q%d" % (FIRST, LAST),
        FormulaRule(formula=['$Q%d="Terminée"' % FIRST], fill=fill(GREEN_BG), font=f(9, True, GREEN)))
    ws.conditional_formatting.add(
        "Q%d:Q%d" % (FIRST, LAST),
        FormulaRule(formula=['$Q%d="En cours"' % FIRST], font=f(9, True, AMBER)))

    note(ws, "A%d" % (LAST + 2),
         "Cotation de la criticité (IPR) :  Gravité — 1 gêne interne … 5 arrêt client ou accident  •  "
         "Fréquence — 1 exceptionnel … 5 quotidien  •  Détection — 1 détecté immédiatement … 5 détecté par le client. "
         "Une action de criticité ≥ 30 doit être traitée en priorité et faire l'objet d'une vérification d'efficacité tracée.",
         "A%d:Y%d" % (LAST + 2, LAST + 3))

    ws.freeze_panes = "E9"
    ws.auto_filter.ref = "A8:Y%d" % LAST
    page(ws, "landscape", title_rows="8:8")
    ws.sheet_view.zoomScale = 85
    return ws
