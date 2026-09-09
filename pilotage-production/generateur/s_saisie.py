# -*- coding: utf-8 -*-
"""Onglet SAISIE_PROD : journal de production et calcul du TRS."""
from openpyxl.styles import Alignment, Font, Border, Side
from openpyxl.worksheet.datavalidation import DataValidation
from openpyxl.formatting.rule import FormulaRule, ColorScaleRule, DataBarRule
from openpyxl.utils import get_column_letter as gcl
from common import *
import data as D

FIRST, LAST = 5, 204          # lignes de donnees

GROUPES = [
    ("IDENTIFICATION", 1, 5, "35566F"),
    ("TEMPS D'OUVERTURE ET PERTES DE TEMPS  (minutes)", 6, 14, "1F4E79"),
    ("PRODUCTION ET QUALITÉ", 15, 20, "2E6B4F"),
    ("RESSOURCES ET SÉCURITÉ", 21, 25, "7A5230"),
    ("ANALYSE DE L'ÉCART", 26, 27, "5B4B7A"),
    ("INDICATEURS CALCULÉS  —  ne rien saisir", 28, 45, "0F2A44"),
]

COLS = [
    # (lettre, en-tete, largeur, type, format)
    ("A", "Date", 11, "in", DATE),
    ("B", "Équipe", 20, "in", None),
    ("C", "Ligne / ressource", 22, "in", None),
    ("D", "N° d'OF / lot", 12, "in", None),
    ("E", "Référence produit", 14, "in", None),
    ("F", "Temps d'ouverture\nTO (min)", 11, "in", NUM),
    ("G", "Arrêts planifiés\n(min)", 10, "in", NUM),
    ("H", "Pannes\n(min)", 9, "in", NUM),
    ("I", "Changement\nde série (min)", 11, "in", NUM),
    ("J", "Réglages et\nmicro-arrêts (min)", 11, "in", NUM),
    ("K", "Manque\nmatière (min)", 10, "in", NUM),
    ("L", "Manque\npersonnel (min)", 10, "in", NUM),
    ("M", "Autres arrêts\nsubis (min)", 10, "in", NUM),
    ("N", "Nombre de\npannes", 9, "in", NUM),
    ("O", "Cadence nominale\n(pcs/min)", 12, "ref", "0.0"),
    ("P", "Qté produite", 11, "in", NUM),
    ("Q", "Qté conforme\n1er passage", 12, "in", NUM),
    ("R", "Qté retouchée", 11, "in", NUM),
    ("S", "Qté rebutée", 10, "in", NUM),
    ("T", "Qté demandée\n(programme)", 12, "in", NUM),
    ("U", "Effectif\nprévu", 9, "in", NUM1),
    ("V", "Effectif\nprésent", 9, "in", NUM1),
    ("W", "Heures\nsupp.", 9, "in", NUM1),
    ("X", "Accidents\navec arrêt", 9, "in", NUM),
    ("Y", "Presqu'accidents\net situations dang.", 11, "in", NUM),
    ("Z", "Code arrêt\nprincipal", 11, "in", None),
    ("AA", "Problème principal du poste", 40, "in", None),
    ("AB", "Temps requis\nTR (min)", 10, "ca", NUM),
    ("AC", "Arrêts subis\n(min)", 10, "ca", NUM),
    ("AD", "Temps de marche\nTF (min)", 11, "ca", NUM),
    ("AE", "Temps utile\nTU (min)", 10, "ca", NUM),
    ("AF", "Disponibilité", 11, "ca", PCT),
    ("AG", "Performance", 11, "ca", PCT),
    ("AH", "Qualité RFT", 10, "ca", PCT),
    ("AI", "TRS", 10, "ca", PCT),
    ("AJ", "TRG", 10, "ca", PCT),
    ("AK", "MTBF (min)", 10, "ca", NUM),
    ("AL", "MTTR (min)", 10, "ca", NUM),
    ("AM", "Non-conformités\n(ppm)", 12, "ca", NUM),
    ("AN", "Taux de service", 11, "ca", PCT),
    ("AO", "Taux de présence", 11, "ca", PCT),
    ("AP", "Perte de cadence\n(min)", 11, "ca", NUM),
    ("AQ", "Perte non-qualité\n(min)", 11, "ca", NUM),
    ("AR", "Contrôle de cohérence", 30, "ca", None),
    ("AS", "Filtre cockpit", 10, "ca", "0"),
]

FORMULES = {
    "O":  '=IF($E{r}="","",IFERROR(INDEX(PARAMETRES!$O$7:$O$36,MATCH($E{r},PARAMETRES!$M$7:$M$36,0)),""))',
    "AB": '=IF($A{r}="","",MAX(0,$F{r}-$G{r}))',
    "AC": '=IF($A{r}="","",SUM($H{r}:$M{r}))',
    "AD": '=IF($A{r}="","",MAX(0,$AB{r}-$AC{r}))',
    "AE": '=IF($A{r}="","",IFERROR($Q{r}/$O{r},""))',
    "AF": '=IF($A{r}="","",IFERROR($AD{r}/$AB{r},""))',
    "AG": '=IF($A{r}="","",IFERROR(($P{r}/$O{r})/$AD{r},""))',
    "AH": '=IF($A{r}="","",IFERROR($Q{r}/$P{r},""))',
    "AI": '=IF($A{r}="","",IFERROR($AF{r}*$AG{r}*$AH{r},""))',
    "AJ": '=IF($A{r}="","",IFERROR($AE{r}/$F{r},""))',
    "AK": '=IF($A{r}="","",IFERROR($AD{r}/$N{r},""))',
    "AL": '=IF($A{r}="","",IFERROR($H{r}/$N{r},""))',
    "AM": '=IF($A{r}="","",IFERROR(($P{r}-$Q{r})/$P{r}*1000000,""))',
    "AN": '=IF($A{r}="","",IFERROR(($Q{r}+$R{r})/$T{r},""))',
    "AO": '=IF($A{r}="","",IFERROR($V{r}/$U{r},""))',
    "AP": '=IF($A{r}="","",IFERROR(MAX(0,$AD{r}-$P{r}/$O{r}),""))',
    "AQ": '=IF($A{r}="","",IFERROR(($P{r}-$Q{r})/$O{r},""))',
    "AR": ('=IF($A{r}="","",'
           'IF($F{r}="","Temps d\'ouverture manquant",'
           'IF($O{r}="","Cadence produit inconnue",'
           'IF($G{r}>$F{r},"Arrêts planifiés > temps d\'ouverture",'
           'IF($AC{r}>$AB{r},"Arrêts subis > temps requis",'
           'IF($P{r}<>$Q{r}+$R{r}+$S{r},"Produite ≠ conforme + retouchée + rebutée",'
           'IF($Q{r}>$P{r},"Conforme > produite",'
           'IF($V{r}>$U{r},"Présents > prévus","OK"))))))))'),
    "AS": ('=IF($A{r}="",0,IF(AND(OR(COCKPIT!$H$4="Toutes",$C{r}=COCKPIT!$H$4),'
           'OR(COCKPIT!$K$4="Toutes",$B{r}=COCKPIT!$K$4)),1,0))'),
}


def build(wb, saisie):
    ws = wb.create_sheet("SAISIE_PROD")
    ws.sheet_view.showGridLines = False
    ws.sheet_properties.tabColor = "1F6FB2"
    ncol = len(COLS)

    title_band(ws, 1, ncol, "SAISIE QUOTIDIENNE DE PRODUCTION",
               "Une ligne par jour, par équipe et par ligne de production. Saisir uniquement les colonnes sur fond jaune "
               "(5 minutes en fin de poste) : les colonnes bleues sont calculées et alimentent le cockpit.")

    for lib, c1, c2, col in GROUPES:
        ws.merge_cells(start_row=3, start_column=c1, end_row=3, end_column=c2)
        c = ws.cell(row=3, column=c1, value=lib)
        c.font = f(9, True, WHITE)
        c.alignment = Alignment(horizontal="center", vertical="center")
        for k in range(c1, c2 + 1):
            ws.cell(row=3, column=k).fill = fill(col)
            ws.cell(row=3, column=k).border = Border(right=Side(style="thin", color=WHITE))
    ws.row_dimensions[3].height = 20

    header_row(ws, 4, [c[1] for c in COLS], height=42)
    for i, (lt, _, w, kind, _) in enumerate(COLS):
        ws.column_dimensions[lt].width = w
        if kind != "in":
            ws.cell(row=4, column=i + 1).fill = fill(NAVY)

    for i, (lt, _, w, kind, fmt) in enumerate(COLS):
        col = i + 1
        rng = "%s%d:%s%d" % (lt, FIRST, lt, LAST)
        if kind == "in":
            apply_range(ws, rng, font=f(9, color="0000FF"), fillc=INPUT_BG, numfmt=fmt,
                        align="center" if fmt else "left")
        else:
            apply_range(ws, rng, font=f(9, color="1A1A1A"), fillc=CALC_BG, numfmt=fmt,
                        align="center")
    apply_range(ws, "AA%d:AA%d" % (FIRST, LAST), font=f(9, color="0000FF"),
                fillc=INPUT_BG, align="left")
    apply_range(ws, "AR%d:AR%d" % (FIRST, LAST), font=f(9), fillc=CALC_BG, align="left")

    for r in range(FIRST, LAST + 1):
        for lt, formule in FORMULES.items():
            ws["%s%d" % (lt, r)] = formule.format(r=r)
        ws.row_dimensions[r].height = 15
    apply_range(ws, "O%d:O%d" % (FIRST, LAST), font=f(9, color=GREEN), fillc=REF_BG,
                numfmt="0.0", align="center")

    # ------------------------------------------------ donnees de demonstration
    cad = {p[0]: p[2] for p in D.PRODUITS}
    for i, s in enumerate(saisie):
        r = FIRST + i
        vals = {"A": s["date"], "B": s["equipe"], "C": s["ligne"], "D": s["of"],
                "E": s["ref"], "F": s["to"], "G": s["planifie"], "H": s["pannes"],
                "I": s["cds"], "J": s["reglages"], "K": s["matiere"], "L": s["personnel"],
                "M": s["autres"], "N": s["n_pannes"], "P": s["produite"], "Q": s["conforme"],
                "R": s["retouche"], "S": s["rebut"], "T": s["demande"], "U": s["eff_prev"],
                "V": s["eff_pres"], "W": s["hsup"], "X": s["acc"], "Y": s["presq"],
                "Z": s["code"], "AA": s["top"]}
        for lt, v in vals.items():
            ws["%s%d" % (lt, r)] = v
        ws["A%d" % r].number_format = DATE

    # ------------------------------------------------ validations
    def dv(kind, rng, **kw):
        v = DataValidation(allow_blank=True, showErrorMessage=True, **kw)
        v.type = kind
        ws.add_data_validation(v)
        v.add(rng)
        return v

    for col, name in [("B", "LST_EQUIPES"), ("C", "LST_LIGNES"), ("E", "LST_PRODUITS"),
                      ("Z", "LST_CAUSES")]:
        dv("list", "%s%d:%s%d" % (col, FIRST, col, LAST), formula1=name,
           errorTitle="Valeur hors référentiel",
           error="Choisir une valeur de la liste (onglet PARAMETRES).")
    dv("date", "A%d:A%d" % (FIRST, LAST), operator="between",
       formula1="DATE(2020,1,1)", formula2="DATE(2099,12,31)",
       errorTitle="Date invalide", error="Saisir une date comprise entre 2020 et 2099.")
    for col in "FGHIJKLMNPQRSTUVWXY":
        dv("decimal", "%s%d:%s%d" % (col, FIRST, col, LAST), operator="greaterThanOrEqual",
           formula1="0", errorTitle="Valeur négative",
           error="Cette colonne n'accepte que des valeurs positives ou nulles.")

    # ------------------------------------------------ mises en forme conditionnelles
    ws.conditional_formatting.add(
        "AI%d:AI%d" % (FIRST, LAST),
        DataBarRule(start_type="num", start_value=0, end_type="num", end_value=1,
                    color="1F6FB2", showValue=True))
    for col in ("AF", "AG", "AH", "AJ"):
        ws.conditional_formatting.add(
            "%s%d:%s%d" % (col, FIRST, col, LAST),
            ColorScaleRule(start_type="num", start_value=0.5, start_color="F8B4B4",
                           mid_type="num", mid_value=0.8, mid_color="FDE9A9",
                           end_type="num", end_value=1.0, end_color="B7E4C0"))
    ws.conditional_formatting.add(
        "A%d:AR%d" % (FIRST, LAST),
        FormulaRule(formula=['AND($AR%d<>"",$AR%d<>"OK")' % (FIRST, FIRST)],
                    fill=fill(RED_BG), font=f(9, True, RED), stopIfTrue=False))
    ws.conditional_formatting.add(
        "X%d:X%d" % (FIRST, LAST),
        FormulaRule(formula=['$X%d>0' % FIRST], fill=fill("F5B7B1"), font=f(9, True, RED)))

    ws.freeze_panes = "F5"
    ws.auto_filter.ref = "A4:%s%d" % (COLS[-1][0], LAST)
    page(ws, "landscape", title_rows="3:4")
    ws.sheet_view.zoomScale = 85
    return ws
