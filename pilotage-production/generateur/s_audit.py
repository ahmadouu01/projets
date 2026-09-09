# -*- coding: utf-8 -*-
"""Onglet AUDIT_5S : grille d'audit ponderee et scoring."""
import datetime as dt
from openpyxl.styles import Alignment, Border, Side
from openpyxl.worksheet.datavalidation import DataValidation
from openpyxl.formatting.rule import FormulaRule, ColorScaleRule, DataBarRule
from openpyxl.chart import RadarChart, LineChart, Reference
from openpyxl.chart.series import SeriesLabel
from common import *
import data as D

FIRST = 8
LAST = FIRST + len(D.AUDIT_5S) - 1        # 32
SC = LAST + 3                              # premiere ligne de scoring
HIST = SC + 9

DOMAINES = ["1S Débarrasser", "2S Ranger", "3S Nettoyer", "4S Standardiser", "5S Progresser"]


def build(wb):
    ws = wb.create_sheet("AUDIT_5S")
    ws.sheet_view.showGridLines = False
    ws.sheet_properties.tabColor = "6B8E23"
    NC = 12

    title_band(ws, 1, NC, "AUDIT 5S  —  grille terrain pondérée",
               "Notation de chaque critère de 0 à 4 :  0 non fait  •  1 amorcé  •  2 partiel  •  3 conforme  •  4 exemplaire et tenu dans la durée. "
               "La pondération traduit l'importance du critère pour la zone auditée (1 = utile, 2 = important, 3 = majeur).")

    band(ws, 4, 1, NC, "IDENTIFICATION DE L'AUDIT")
    ident = [("A5", "Zone auditée :", "C5", "Ligne 1 - Assemblage"),
             ("E5", "Date de l'audit :", "F5", dt.date(2026, 1, 30)),
             ("H5", "Auditeur :", "I5", "LEROY S."),
             ("A6", "Zone de référence :", "C6", "UAP Assemblage-Usinage"),
             ("E6", "Audit précédent :", "F6", dt.date(2025, 12, 19)),
             ("H6", "Accompagné de :", "I6", "Équipe A (matin)")]
    for lab_ref, lab, val_ref, val in ident:
        label(ws, lab_ref, lab, bold=True, size=9, align="right")
        c = ws[val_ref]
        c.value = val
        c.font = f(9, True, "0000FF")
        c.fill = fill(INPUT_BG)
        c.border = BOX
        c.alignment = Alignment(horizontal="left", vertical="center", indent=1)
        if isinstance(val, dt.date):
            c.number_format = DATE
    for rng in ("C5:D5", "F5:G5", "I5:L5", "C6:D6", "F6:G6", "I6:L6"):
        ws.merge_cells(rng)
        for row in ws[rng]:
            for cc in row:
                cc.fill = fill(INPUT_BG)
                cc.border = BOX

    entetes = ["Domaine", "N°", "Critère observé sur le terrain", "Pondération",
               "Note\n(0 à 4)", "Note\npondérée", "Note\naudit N-1", "Écart",
               "Constat terrain", "Action corrective décidée", "Pilote", "Délai"]
    header_row(ws, FIRST - 1, entetes, height=40)

    for i, (dom, num, crit, pond, note_n, note_n1) in enumerate(D.AUDIT_5S):
        r = FIRST + i
        ws["A%d" % r] = dom
        ws["B%d" % r] = num
        ws["C%d" % r] = crit
        ws["D%d" % r] = pond
        ws["E%d" % r] = note_n
        ws["G%d" % r] = note_n1
        ws["F%d" % r] = '=IF($E{r}="","",$D{r}*$E{r})'.format(r=r)
        ws["H%d" % r] = '=IF(OR($E{r}="",$G{r}=""),"",$E{r}-$G{r})'.format(r=r)
        for col in range(1, 13):
            c = ws.cell(row=r, column=col)
            c.border = BOX
            c.font = f(9)
            c.alignment = Alignment(horizontal="left", vertical="center", indent=1, wrap_text=True)
        for col, fmt in ((2, "0"), (4, "0"), (5, "0"), (6, "0"), (7, "0"), (8, "+0;-0;0")):
            c = ws.cell(row=r, column=col)
            c.alignment = Alignment(horizontal="center", vertical="center")
            c.number_format = fmt
        for col in (4, 5, 7):
            c = ws.cell(row=r, column=col)
            c.font = f(9, True, "0000FF")
            c.fill = fill(INPUT_BG)
        for col in (6, 8):
            ws.cell(row=r, column=col).fill = fill(CALC_BG)
        for col in (9, 10, 11, 12):
            ws.cell(row=r, column=col).font = f(9, color="0000FF")
            ws.cell(row=r, column=col).fill = fill(INPUT_BG)
        ws.cell(row=r, column=1).font = f(9, True, STEEL)
        ws.cell(row=r, column=12).number_format = DATE
        ws.row_dimensions[r].height = 22

    # ---------------------------------------------------------- scoring
    band(ws, SC - 2, 1, NC, "RÉSULTAT DE L'AUDIT  —  score pondéré par domaine")
    header_row(ws, SC - 1, ["Domaine 5S", "Note pondérée\nobtenue", "Note pondérée\nmaximale",
                            "Score", "Cible", "Écart", "Statut", "Score audit N-1",
                            "Tendance", "", "", ""], height=34)
    cible = '=INDEX(PARAMETRES!$D$16:$D$29,MATCH("AUDIT5S",PARAMETRES!$A$16:$A$29,0))'
    for i, dom in enumerate(DOMAINES + ["SCORE GLOBAL"]):
        r = SC + i
        glob = (i == len(DOMAINES))
        ws["A%d" % r] = dom
        if glob:
            ws["B%d" % r] = '=SUM($B${a}:$B${b})'.format(a=SC, b=SC + 4)
            ws["C%d" % r] = '=SUM($C${a}:$C${b})'.format(a=SC, b=SC + 4)
            ws["H%d" % r] = '=IFERROR(SUMPRODUCT($D${f}:$D${l},$G${f}:$G${l})/($C{r}),"")'.format(
                f=FIRST, l=LAST, r=r)
        else:
            ws["B%d" % r] = '=SUMIF($A${f}:$A${l},$A{r},$F${f}:$F${l})'.format(f=FIRST, l=LAST, r=r)
            ws["C%d" % r] = '=SUMIF($A${f}:$A${l},$A{r},$D${f}:$D${l})*4'.format(f=FIRST, l=LAST, r=r)
            ws["H%d" % r] = ('=IFERROR(SUMPRODUCT(($A${f}:$A${l}=$A{r})*$D${f}:$D${l}*$G${f}:$G${l})'
                             '/$C{r},"")').format(f=FIRST, l=LAST, r=r)
        ws["D%d" % r] = '=IFERROR($B{r}/$C{r},"")'.format(r=r)
        ws["E%d" % r] = cible
        ws["F%d" % r] = '=IF($D{r}="","",$D{r}-$E{r})'.format(r=r)
        ws["G%d" % r] = ('=IF($D{r}="","",IF($D{r}>=$E{r},"Conforme",'
                         'IF($D{r}>=$E{r}-0.1,"À consolider","Non conforme")))').format(r=r)
        ws["I%d" % r] = ('=IF(OR($D{r}="",$H{r}=""),"",IF($D{r}>$H{r}+0.02,"En progrès",'
                         'IF($D{r}<$H{r}-0.02,"En recul","Stable")))').format(r=r)
        for col in range(1, 10):
            c = ws.cell(row=r, column=col)
            c.border = BOX
            c.font = f(9, glob)
            c.fill = fill(LIGHT if glob else CALC_BG)
            c.alignment = Alignment(horizontal="center", vertical="center")
        ws.cell(row=r, column=1).alignment = Alignment(horizontal="left", vertical="center", indent=1)
        ws.cell(row=r, column=1).font = f(9.5, True, NAVY if glob else STEEL)
        for col, fmt in ((2, "0"), (3, "0"), (4, PCT), (5, PCT), (6, "+0.0%;-0.0%;0.0%"), (8, PCT)):
            ws.cell(row=r, column=col).number_format = fmt
        ws.row_dimensions[r].height = 18

    for rng, ok, ko in (("G%d:G%d" % (SC, SC + 5), "Conforme", "Non conforme"),):
        ws.conditional_formatting.add(rng, FormulaRule(formula=['$G%d="Conforme"' % SC],
                                      fill=fill(GREEN_BG), font=f(9, True, GREEN)))
        ws.conditional_formatting.add(rng, FormulaRule(formula=['$G%d="Non conforme"' % SC],
                                      fill=fill(RED_BG), font=f(9, True, RED)))
        ws.conditional_formatting.add(rng, FormulaRule(formula=['$G%d="À consolider"' % SC],
                                      fill=fill(AMBER_BG), font=f(9, True, AMBER)))
    ws.conditional_formatting.add(
        "D%d:D%d" % (SC, SC + 5),
        DataBarRule(start_type="num", start_value=0, end_type="num", end_value=1,
                    color="6B8E23", showValue=True))
    ws.conditional_formatting.add(
        "E%d:E%d" % (FIRST, LAST),
        ColorScaleRule(start_type="num", start_value=0, start_color="F5A6A6",
                       mid_type="num", mid_value=2, mid_color="FDE9A9",
                       end_type="num", end_value=4, end_color="B7E4C0"))
    ws.conditional_formatting.add(
        "H%d:H%d" % (FIRST, LAST),
        FormulaRule(formula=['AND($H%d<>"",$H%d<0)' % (FIRST, FIRST)], font=f(9, True, RED)))

    # ---------------------------------------------------------- historique
    band(ws, HIST - 1, 1, NC, "HISTORIQUE DES AUDITS DE LA ZONE  —  à compléter à chaque passage")
    header_row(ws, HIST, ["Date de l'audit", "Zone", "Score global", "Actions ouvertes",
                          "Actions soldées", "Auditeur", "", "", "", "", "", ""], height=26)
    hist = [(dt.date(2025, 9, 26), "Ligne 1 - Assemblage", 0.52, 8, 5, "LEROY S."),
            (dt.date(2025, 10, 24), "Ligne 1 - Assemblage", 0.58, 7, 6, "LEROY S."),
            (dt.date(2025, 11, 21), "Ligne 1 - Assemblage", 0.61, 6, 5, "DURAND M."),
            (dt.date(2025, 12, 19), "Ligne 1 - Assemblage", 0.63, 6, 4, "LEROY S.")]
    for i in range(12):
        r = HIST + 1 + i
        for col in range(1, 7):
            c = ws.cell(row=r, column=col)
            c.border = BOX
            c.font = f(9, color="0000FF")
            c.fill = fill(INPUT_BG)
            c.alignment = Alignment(horizontal="center", vertical="center")
        ws.cell(row=r, column=1).number_format = DATE
        ws.cell(row=r, column=3).number_format = PCT
        ws.cell(row=r, column=2).alignment = Alignment(horizontal="left", vertical="center", indent=1)
        if i < len(hist):
            for j, v in enumerate(hist[i]):
                ws.cell(row=r, column=1 + j, value=v)
        ws.row_dimensions[r].height = 15
    r = HIST + 1 + len(hist)
    ws.cell(row=r, column=1, value='=$F$5')
    ws.cell(row=r, column=1).number_format = DATE
    ws.cell(row=r, column=2, value='=$C$5')
    ws.cell(row=r, column=3, value='=$D${r}'.format(r=SC + 5))
    for col in (1, 2, 3):
        ws.cell(row=r, column=col).font = f(9, True, GREEN)
        ws.cell(row=r, column=col).fill = fill(CALC_BG)

    # ---------------------------------------------------------- graphiques
    radar = RadarChart()
    radar.type = "marker"
    radar.style = 26
    radar.title = "Profil 5S de la zone — score par domaine"
    radar.height, radar.width = 9.5, 12
    radar.add_data(Reference(ws, min_col=4, min_row=SC, max_row=SC + 4), titles_from_data=False)
    radar.add_data(Reference(ws, min_col=5, min_row=SC, max_row=SC + 4), titles_from_data=False)
    radar.set_categories(Reference(ws, min_col=1, min_row=SC, max_row=SC + 4))
    radar.series[0].tx = SeriesLabel(v="Score de l'audit")
    radar.series[1].tx = SeriesLabel(v="Cible")
    radar.y_axis.scaling.min = 0
    radar.y_axis.scaling.max = 1
    str_categories(radar, "'AUDIT_5S'!$A$%d:$A$%d" % (SC, SC + 4))
    ws.add_chart(radar, "A%d" % (HIST + 15))

    ln = LineChart()
    ln.style = 2
    ln.title = "Évolution du score 5S de la zone"
    ln.height, ln.width = 9.5, 14
    ln.add_data(Reference(ws, min_col=3, min_row=HIST + 1, max_row=HIST + 12), titles_from_data=False)
    ln.set_categories(Reference(ws, min_col=1, min_row=HIST + 1, max_row=HIST + 12))
    ln.series[0].tx = SeriesLabel(v="Score global")
    ln.y_axis.numFmt = "0%"
    ln.y_axis.majorGridlines = None
    ws.add_chart(ln, "F%d" % (HIST + 15))

    dv = DataValidation(type="list", formula1='"0,1,2,3,4"', allow_blank=True,
                        showErrorMessage=True, errorTitle="Notation 5S",
                        error="Noter de 0 à 4.")
    ws.add_data_validation(dv)
    dv.add("E%d:E%d" % (FIRST, LAST))
    dv.add("G%d:G%d" % (FIRST, LAST))
    dv2 = DataValidation(type="whole", operator="between", formula1="1", formula2="3",
                         allow_blank=True, showErrorMessage=True, errorTitle="Pondération",
                         error="Pondérer de 1 à 3.")
    ws.add_data_validation(dv2)
    dv2.add("D%d:D%d" % (FIRST, LAST))
    dv3 = DataValidation(type="list", formula1="LST_PILOTES", allow_blank=True)
    ws.add_data_validation(dv3)
    dv3.add("K%d:K%d" % (FIRST, LAST))

    widths(ws, {"A": 16, "B": 5, "C": 52, "D": 11, "E": 8, "F": 9, "G": 9, "H": 8,
                "I": 34, "J": 38, "K": 14, "L": 12})
    ws.freeze_panes = "A8"
    page(ws, "landscape", title_rows="7:7")
    return ws
