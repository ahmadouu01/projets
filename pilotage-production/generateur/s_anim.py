# -*- coding: utf-8 -*-
"""Onglet ANIMATION : board de management visuel du point 5 minutes (A3 paysage)."""
from openpyxl.styles import Alignment, Border, Side, Font
from openpyxl.worksheet.datavalidation import DataValidation
from openpyxl.formatting.rule import FormulaRule, DataBarRule
from openpyxl.chart import DoughnutChart, BarChart, LineChart, Reference
from openpyxl.chart.series import SeriesLabel
from openpyxl.chart.shapes import GraphicalProperties
from openpyxl.chart.marker import DataPoint, Marker
from openpyxl.chart.label import DataLabelList
from openpyxl.drawing.line import LineProperties
from openpyxl.utils import get_column_letter as gcl
from common import *
import s_calc as C

NC = 46                      # colonnes du board
D1, D31 = 8, 38              # colonnes des 31 jours
GRID = 8                     # premiere ligne de la grille SQCDP

AXES = [
    ("S", "Sécurité",          "W", "Accident ou situation dangereuse"),
    ("Q", "Qualité",           "X", "Taux de conformité au premier passage"),
    ("C", "Coût / rendement",  "Y", "TRS de la journée"),
    ("D", "Délai / service",   "Z", "Respect du programme"),
    ("P", "Personnel",         "AA", "Effectif présent contre effectif prévu"),
]

VERT, ORANGE, ROUGE, VIDE = "2E9E5B", "E8A33D", "C0392B", "E9EDF1"


def build(wb):
    ws = wb.create_sheet("ANIMATION")
    ws.sheet_view.showGridLines = False
    ws.sheet_properties.tabColor = "C0392B"
    for col in range(1, NC + 1):
        ws.column_dimensions[gcl(col)].width = 3.3

    title_band(ws, 1, NC, "MANAGEMENT VISUEL  —  POINT 5 MINUTES",
               '=PARAMETRES!$C$5&"   |   "&PARAMETRES!$C$6&"   |   "&COCKPIT!$E$4&'
               '"   |   Édition du "&TEXT(TODAY(),"dd/mm/yyyy")', height=38)
    ws["A1"].font = Font(name=FONT, size=19, bold=True, color=WHITE)
    ws["A2"].font = f(10, False, "4A5A6A", italic=True)

    # ---------------------------------------------------------- bandeau contexte
    ctx = [("Périmètre", '=COCKPIT!$H$4', 1, 12),
           ("Équipe", '=COCKPIT!$K$4', 13, 23),
           ("Jours de production saisis", '=CALC!$B$9', 24, 34),
           ("Actions ouvertes en retard", '=PLAN_ACTIONS!$D$6', 35, 46)]
    for lib, formule, c1, c2 in ctx:
        ws.merge_cells(start_row=4, start_column=c1, end_row=4, end_column=c2)
        c = ws.cell(row=4, column=c1, value=lib)
        c.font = f(8, True, GREY)
        c.alignment = Alignment(horizontal="center", vertical="bottom")
        ws.merge_cells(start_row=5, start_column=c1, end_row=5, end_column=c2)
        v = ws.cell(row=5, column=c1, value=formule)
        v.font = f(12, True, NAVY)
        v.alignment = Alignment(horizontal="center", vertical="center")
        for k in range(c1, c2 + 1):
            ws.cell(row=4, column=k).fill = fill(LIGHT)
            ws.cell(row=5, column=k).fill = fill(LIGHT)
        ws.cell(row=5, column=c2).border = Border(right=Side(style="thin", color=WHITE))
    ws.row_dimensions[4].height = 14
    ws.row_dimensions[5].height = 22

    # ---------------------------------------------------------- grille SQCDP
    band(ws, GRID - 2, 1, NC,
         "1 ·  S Q C D P   —   LE MOIS JOUR PAR JOUR      "
         "(vert : objectif tenu   ·   orange : vigilance   ·   rouge : écart à traiter)",
         tone=NAVY, color=WHITE, size=11)
    for k in range(1, NC + 1):
        ws.cell(row=GRID - 2, column=k).fill = fill(NAVY)

    ws.merge_cells(start_row=GRID - 1, start_column=1, end_row=GRID - 1, end_column=D1 - 1)
    c = ws.cell(row=GRID - 1, column=1, value="Axe de performance")
    c.font = f(8.5, True, STEEL)
    c.alignment = Alignment(horizontal="center", vertical="center")
    for j in range(31):
        c = ws.cell(row=GRID - 1, column=D1 + j, value=j + 1)
        c.font = f(7.5, True, GREY)
        c.alignment = Alignment(horizontal="center", vertical="center")
        c.number_format = "0"
    for lib, c1, c2 in (("Jours au vert", 39, 41), ("Vigilance", 42, 43), ("Écarts", 44, 46)):
        ws.merge_cells(start_row=GRID - 1, start_column=c1, end_row=GRID - 1, end_column=c2)
        c = ws.cell(row=GRID - 1, column=c1, value=lib)
        c.font = f(7.5, True, STEEL)
        c.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
    ws.row_dimensions[GRID - 1].height = 22

    thinw = Side(style="thin", color=WHITE)
    for i, (lettre, nom, col_calc, _desc) in enumerate(AXES):
        r = GRID + i
        ws.merge_cells(start_row=r, start_column=1, end_row=r, end_column=2)
        c = ws.cell(row=r, column=1, value=lettre)
        c.font = Font(name=FONT, size=16, bold=True, color=WHITE)
        c.fill = fill(NAVY)
        c.alignment = Alignment(horizontal="center", vertical="center")
        ws.cell(row=r, column=2).fill = fill(NAVY)
        ws.merge_cells(start_row=r, start_column=3, end_row=r, end_column=D1 - 1)
        c = ws.cell(row=r, column=3, value=nom)
        c.font = f(10, True, NAVY)
        c.alignment = Alignment(horizontal="left", vertical="center", indent=1)
        for k in range(3, D1):
            ws.cell(row=r, column=k).fill = fill(LIGHT)

        for j in range(31):
            col = D1 + j
            cc = ws.cell(row=r, column=col,
                         value='=IFERROR(INDEX(CALC!$W$%d:$AA$%d,%d,%d),"")'
                               % (C.JOUR0, C.JOURN, j + 1, i + 1))
            cc.number_format = ";;;"          # la valeur sert uniquement à colorer la case
            cc.fill = fill(VIDE)
            cc.border = Border(left=thinw, right=thinw, top=thinw, bottom=thinw)
        plage = "%s%d:%s%d" % (gcl(D1), r, gcl(D31), r)
        for val, couleur in ((1, VERT), (2, ORANGE), (3, ROUGE)):
            ws.conditional_formatting.add(
                plage, FormulaRule(formula=['%s%d=%d' % (gcl(D1), r, val)], fill=fill(couleur)))
        for lib, c1, c2, val, couleur in (("v", 39, 41, 1, VERT), ("o", 42, 43, 2, ORANGE),
                                          ("r", 44, 46, 3, ROUGE)):
            ws.merge_cells(start_row=r, start_column=c1, end_row=r, end_column=c2)
            cc = ws.cell(row=r, column=c1, value='=COUNTIF(%s,%d)' % (plage, val))
            cc.font = f(12, True, couleur)
            cc.alignment = Alignment(horizontal="center", vertical="center")
            cc.number_format = "0"
            for k in range(c1, c2 + 1):
                ws.cell(row=r, column=k).fill = fill(WHITE)
                ws.cell(row=r, column=k).border = Border(bottom=Side(style="thin", color=BORDERCOL))
        ws.row_dimensions[r].height = 24

    r = GRID + len(AXES)
    ws.merge_cells(start_row=r, start_column=1, end_row=r, end_column=NC)
    note(ws, "A%d" % r,
         "Lecture :  S — accident ou situation dangereuse du jour   ·   Q — qualité au premier passage   ·   "
         "C — TRS de la journée   ·   D — respect du programme   ·   P — effectif présent contre effectif prévu.      "
         "Une case grise signifie qu'aucune production n'a été saisie ce jour-là.")
    ws.row_dimensions[r].height = 16

    # ---------------------------------------------------------- jauges + sécurité
    r0 = r + 2
    band(ws, r0, 1, NC, "2 ·  LES INDICATEURS DU MOIS", tone=NAVY, color=WHITE, size=11)
    for k in range(1, NC + 1):
        ws.cell(row=r0, column=k).fill = fill(NAVY)

    blocs = [(1, 11, 0, "TRS"), (12, 22, 1, "QUAL"), (23, 33, 2, "SERVICE")]
    for c1, c2, idx, code in blocs:
        _jauge(ws, wb, c1, r0 + 1, C.GAU0 + idx)
        rv = r0 + 12
        ws.merge_cells(start_row=rv, start_column=c1, end_row=rv, end_column=c2)
        v = ws.cell(row=rv, column=c1, value='=CALC!$B$%d' % (C.GAU0 + idx))
        v.font = Font(name=FONT, size=24, bold=True, color=NAVY)
        v.number_format = PCT
        v.alignment = Alignment(horizontal="center", vertical="center")
        ws.merge_cells(start_row=rv + 1, start_column=c1, end_row=rv + 1, end_column=c2)
        s = ws.cell(row=rv + 1, column=c1,
                    value='="cible "&TEXT(CALC!$E$%d,"0.0%%")' % (C.GAU0 + idx))
        s.font = f(9, False, GREY)
        s.alignment = Alignment(horizontal="center", vertical="center")
        for val, coul, bg in ((1, GREEN, GREEN_BG), (2, AMBER, AMBER_BG), (3, RED, RED_BG)):
            ws.conditional_formatting.add(
                "%s%d:%s%d" % (gcl(c1), rv, gcl(c2), rv),
                FormulaRule(formula=['COCKPIT!$Y$%d=%d' % (6 + idx * 3, val)],
                            font=Font(name=FONT, size=24, bold=True, color=coul), fill=fill(bg)))
        ws.row_dimensions[rv].height = 32
        ws.row_dimensions[rv + 1].height = 16

    # tuile sécurité
    c1, c2 = 34, NC
    ws.merge_cells(start_row=r0 + 1, start_column=c1, end_row=r0 + 2, end_column=c2)
    t = ws.cell(row=r0 + 1, column=c1, value="SÉCURITÉ")
    t.font = f(11, True, WHITE)
    t.alignment = Alignment(horizontal="center", vertical="center")
    for rr in (r0 + 1, r0 + 2):
        for k in range(c1, c2 + 1):
            ws.cell(row=rr, column=k).fill = fill("C0392B")
    ws.merge_cells(start_row=r0 + 3, start_column=c1, end_row=r0 + 4, end_column=c2)
    lab = ws.cell(row=r0 + 3, column=c1, value="Jours sans accident avec arrêt")
    lab.font = f(9.5, True, "7A2A22")
    lab.alignment = Alignment(horizontal="center", vertical="center")
    ws.merge_cells(start_row=r0 + 5, start_column=c1, end_row=r0 + 8, end_column=c2)
    big = ws.cell(row=r0 + 5, column=c1,
                  value=('=IF(COUNTIF(SAISIE_PROD!$X$5:$X$204,">0")=0,'
                         'COUNT(SAISIE_PROD!$A$5:$A$204),'
                         'TODAY()-SUMPRODUCT(MAX((SAISIE_PROD!$X$5:$X$204>0)*'
                         'SAISIE_PROD!$A$5:$A$204)))'))
    big.font = Font(name=FONT, size=44, bold=True, color="C0392B")
    big.number_format = "0"
    big.alignment = Alignment(horizontal="center", vertical="center")
    ws.merge_cells(start_row=r0 + 9, start_column=c1, end_row=r0 + 9, end_column=c2)
    sub = ws.cell(row=r0 + 9, column=c1,
                  value=('="Accidents du mois : "&CALC!$AC$26&"      Presqu\'accidents : "'
                         '&SUM(CALC!$V$%d:$V$%d)' % (C.JOUR0, C.JOURN)))
    sub.font = f(9.5, True, "7A2A22")
    sub.alignment = Alignment(horizontal="center", vertical="center")
    ws.merge_cells(start_row=r0 + 10, start_column=c1, end_row=r0 + 13, end_column=c2)
    msg = ws.cell(row=r0 + 10, column=c1,
                  value=('=IF(CALC!$AC$26>0,"Un accident est survenu ce mois-ci : '
                         'l\'analyse doit être ouverte et diffusée aux trois équipes.",'
                         '"Aucun accident ce mois-ci. Maintenir la vigilance sur les '
                         'presqu\'accidents et les situations dangereuses.")'))
    msg.font = f(9.5, False, "7A2A22")
    msg.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
    for rr in range(r0 + 3, r0 + 14):
        for k in range(c1, c2 + 1):
            ws.cell(row=rr, column=k).fill = fill("FDEDEC")
    for rr in range(r0 + 1, r0 + 14):
        ws.cell(row=rr, column=c1).border = Border(left=Side(style="medium", color="C0392B"))
        ws.cell(row=rr, column=c2).border = Border(right=Side(style="medium", color="C0392B"))

    # ---------------------------------------------------------- graphiques
    r1 = r0 + 15
    band(ws, r1, 1, NC, "3 ·  OÙ PART LE TEMPS  ·  COMMENT ÉVOLUE LE RENDEMENT",
         tone=NAVY, color=WHITE, size=11)
    for k in range(1, NC + 1):
        ws.cell(row=r1, column=k).fill = fill(NAVY)
    _chart_pertes(ws, wb, "A%d" % (r1 + 1))
    _chart_tendance(ws, wb, "%s%d" % (gcl(24), r1 + 1))

    # ---------------------------------------------------------- actions
    r2 = r1 + 20
    band(ws, r2, 1, NC, "4 ·  LES TROIS ACTIONS PRIORITAIRES", tone=NAVY, color=WHITE, size=11)
    for k in range(1, NC + 1):
        ws.cell(row=r2, column=k).fill = fill(NAVY)
    colonnes = [("N°", 1, 3, "B", None), ("Problème constaté", 4, 18, "C", None),
                ("Action décidée", 19, 32, "I", None), ("Pilote", 33, 37, "D", None),
                ("Échéance", 38, 41, "E", DATE), ("Avancement", 42, 44, "F", PCT),
                ("Retard", 45, 46, "H", "0")]
    for lib, c1, c2, _src, _fmt in colonnes:
        ws.merge_cells(start_row=r2 + 1, start_column=c1, end_row=r2 + 1, end_column=c2)
        c = ws.cell(row=r2 + 1, column=c1, value=lib)
        c.font = f(8.5, True, WHITE)
        c.alignment = Alignment(horizontal="center", vertical="center")
        for k in range(c1, c2 + 1):
            ws.cell(row=r2 + 1, column=k).fill = fill(STEEL)
    ws.row_dimensions[r2 + 1].height = 20
    for i in range(3):
        r = r2 + 2 + i
        for lib, c1, c2, src, fmt in colonnes:
            ws.merge_cells(start_row=r, start_column=c1, end_row=r, end_column=c2)
            c = ws.cell(row=r, column=c1, value='=CALC!$%s$%d' % (src, C.TR3 + i))
            c.font = f(9.5, src == "B")
            c.number_format = fmt or "General"
            c.alignment = Alignment(horizontal="left" if fmt is None else "center",
                                    vertical="center", indent=1, wrap_text=True)
            for k in range(c1, c2 + 1):
                ws.cell(row=r, column=k).fill = fill(WHITE if i % 2 == 0 else ROWALT)
                ws.cell(row=r, column=k).border = Border(bottom=Side(style="thin", color=BORDERCOL))
        ws.row_dimensions[r].height = 30
    ws.conditional_formatting.add(
        "%s%d:%s%d" % (gcl(42), r2 + 2, gcl(42), r2 + 4),
        DataBarRule(start_type="num", start_value=0, end_type="num", end_value=1,
                    color="2E9E5B", showValue=True))
    ws.conditional_formatting.add(
        "%s%d:%s%d" % (gcl(45), r2 + 2, gcl(45), r2 + 4),
        FormulaRule(formula=['AND(%s%d<>"",%s%d>0)' % (gcl(45), r2 + 2, gcl(45), r2 + 2)],
                    fill=fill(RED_BG), font=f(9.5, True, RED)))

    # ---------------------------------------------------------- décisions
    r3 = r2 + 6
    band(ws, r3, 1, NC, "5 ·  DÉCISIONS DU POINT 5 MINUTES  —  à renseigner devant l'équipe",
         tone=NAVY, color=WHITE, size=11)
    for k in range(1, NC + 1):
        ws.cell(row=r3, column=k).fill = fill(NAVY)
    dcol = [("Écart constaté aujourd'hui", 1, 18), ("Décision prise", 19, 34),
            ("Pilote", 35, 41), ("Échéance", 42, 46)]
    for lib, c1, c2 in dcol:
        ws.merge_cells(start_row=r3 + 1, start_column=c1, end_row=r3 + 1, end_column=c2)
        c = ws.cell(row=r3 + 1, column=c1, value=lib)
        c.font = f(8.5, True, WHITE)
        c.alignment = Alignment(horizontal="center", vertical="center")
        for k in range(c1, c2 + 1):
            ws.cell(row=r3 + 1, column=k).fill = fill(SLATE)
    for i in range(3):
        r = r3 + 2 + i
        for lib, c1, c2 in dcol:
            ws.merge_cells(start_row=r, start_column=c1, end_row=r, end_column=c2)
            c = ws.cell(row=r, column=c1)
            c.font = f(9.5, color="0000FF")
            c.alignment = Alignment(horizontal="left", vertical="center", indent=1)
            for k in range(c1, c2 + 1):
                ws.cell(row=r, column=k).fill = fill(INPUT_BG)
                ws.cell(row=r, column=k).border = BOX
        ws.cell(row=r, column=42).number_format = DATE
        ws.row_dimensions[r].height = 24
    dv = DataValidation(type="list", formula1="LST_PILOTES", allow_blank=True)
    ws.add_data_validation(dv)
    dv.add("%s%d:%s%d" % (gcl(35), r3 + 2, gcl(35), r3 + 4))

    ws.page_setup.orientation = "landscape"
    ws.page_setup.paperSize = 8            # A3
    from openpyxl.worksheet.properties import PageSetupProperties
    ws.sheet_properties.pageSetUpPr = PageSetupProperties(fitToPage=True)
    ws.page_setup.fitToWidth = 1
    ws.page_setup.fitToHeight = 1
    ws.print_options.horizontalCentered = True
    ws.page_margins.left = ws.page_margins.right = 0.3
    ws.page_margins.top = ws.page_margins.bottom = 0.3
    ws.print_area = "A1:%s%d" % (gcl(NC), r3 + 5)
    ws.sheet_view.zoomScale = 75
    return ws


def _jauge(ws, wb, col, row, ligne_calc):
    calc = wb["CALC"]
    dn = DoughnutChart(holeSize=64, firstSliceAng=270)
    dn.add_data(Reference(calc, min_col=2, max_col=4, min_row=ligne_calc, max_row=ligne_calc),
                from_rows=True, titles_from_data=False)
    dn.series[0].tx = SeriesLabel(v="Jauge")
    dn.series[0].data_points = [
        DataPoint(idx=0, spPr=GraphicalProperties(solidFill="1F6FB2",
                                                  ln=LineProperties(noFill=True))),
        DataPoint(idx=1, spPr=GraphicalProperties(solidFill="DDE4EB",
                                                  ln=LineProperties(noFill=True))),
        DataPoint(idx=2, spPr=GraphicalProperties(noFill=True,
                                                  ln=LineProperties(noFill=True)))]
    dn.title = None
    dn.legend = None
    dn.height, dn.width = 5.6, 5.9
    dn.graphical_properties = GraphicalProperties(noFill=True, ln=LineProperties(noFill=True))
    ws.add_chart(dn, "%s%d" % (gcl(col + 1), row))


def _chart_pertes(ws, wb, anchor):
    calc = wb["CALC"]
    b = BarChart()
    b.type = "bar"
    b.title = "Où part le temps — minutes perdues sur le mois"
    b.height, b.width = 9.6, 13.9
    b.add_data(Reference(calc, min_col=7, min_row=C.PAR0, max_row=C.PARN), titles_from_data=False)
    b.set_categories(Reference(calc, min_col=6, min_row=C.PAR0, max_row=C.PARN))
    str_categories(b, "'CALC'!$F$%d:$F$%d" % (C.PAR0, C.PARN))
    b.series[0].tx = SeriesLabel(v="Minutes perdues")
    b.series[0].graphicalProperties = GraphicalProperties(solidFill="1F6FB2")
    b.gapWidth = 40
    b.legend = None
    b.dLbls = DataLabelList()
    b.dLbls.showVal = True
    b.x_axis.majorGridlines = None
    ws.add_chart(b, anchor)


def _chart_tendance(ws, wb, anchor):
    calc = wb["CALC"]
    ln = LineChart()
    ln.title = "Rendement jour par jour — cible et limites de variation courante"
    ln.height, ln.width = 9.6, 13.9
    for idx in (18, 19, 14, 16, 17):
        ln.add_data(Reference(calc, min_col=idx, min_row=C.JOUR0, max_row=C.JOURN),
                    titles_from_data=False)
    ln.set_categories(Reference(calc, min_col=1, min_row=C.JOUR0, max_row=C.JOURN))
    styles = [("TRS du jour", "1F6FB2", None, 24000, True),
              ("Moyenne mobile 7 jours", "B5651D", None, 22000, False),
              ("Cible", "1E7B34", "dash", 15000, False),
              ("Limite haute", "9AA7B2", "sysDot", 10000, False),
              ("Limite basse", "9AA7B2", "sysDot", 10000, False)]
    for i, (lib, color, dash, w, mark) in enumerate(styles):
        s = ln.series[i]
        s.tx = SeriesLabel(v=lib)
        lp = LineProperties(solidFill=color, w=w)
        if dash:
            lp.prstDash = dash
        s.graphicalProperties = GraphicalProperties(ln=lp)
        s.marker = Marker(symbol="circle", size=4) if mark else Marker(symbol="none")
        s.smooth = False
    ln.y_axis.numFmt = "0%"
    ln.y_axis.majorGridlines = None
    ln.x_axis.title = "Jour du mois"
    ws.add_chart(ln, anchor)
