# -*- coding: utf-8 -*-
"""Génère trois extractions ERP réalistes, prêtes à être importées dans le dashboard.

Les données des trois fichiers sont cohérentes entre elles :
les arrêts du journal alimentent exactement les rubriques de la production
journalière, et les ordres de fabrication reprennent les mêmes lignes,
articles et cadences.
"""
import csv, datetime as dt, os, random
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment
from openpyxl.utils import get_column_letter as gcl

random.seed(70712)
OUT = os.path.join(os.path.dirname(__file__), "..", "exemples-erp")
SITE = "Usine de Vernon"
DEB, FIN = dt.date(2025, 10, 1), dt.date(2026, 1, 30)
FERIES = {dt.date(2025, 11, 1), dt.date(2025, 11, 11), dt.date(2025, 12, 25),
          dt.date(2026, 1, 1)}
CONGES = (dt.date(2025, 12, 22), dt.date(2026, 1, 2))     # arrêt annuel

LIGNES = ["LA-01 Assemblage", "LU-02 Usinage", "LC-03 Conditionnement"]
EQUIPES = {"LA-01 Assemblage": ["Matin", "Après-midi"],
           "LU-02 Usinage": ["Matin", "Après-midi", "Nuit"],
           "LC-03 Conditionnement": ["Matin"]}
ARTICLES = [
    ("ART-40120", "Carter aluminium A20", "LA-01 Assemblage", 15.0),
    ("ART-40130", "Carter aluminium A30", "LA-01 Assemblage", 12.5),
    ("ART-51210", "Support moteur S12", "LU-02 Usinage", 9.0),
    ("ART-51810", "Support moteur S18", "LU-02 Usinage", 7.5),
    ("ART-62400", "Kit conditionné K4", "LC-03 Conditionnement", 22.0),
    ("ART-62600", "Kit conditionné K6", "LC-03 Conditionnement", 18.0),
]
CLIENTS = ["AUTOMOTIVE SA", "NORDIC MACHINES", "GROUPE VALTEC", "SIDER INDUSTRIE",
           "MECAPLUS", "ATLANTIQUE EQUIP.", "TECHNOFORGE"]

# code motif : (libellé, rubrique de saisie, minutes de base)
MOTIFS = {
    "PL-10": ("Pause légale et relève de poste", "plan", 40),
    "PL-20": ("Maintenance préventive planifiée", "plan", 55),
    "PL-30": ("Réunion, formation, audit", "plan", 45),
    "MA-10": ("Panne mécanique", "pannes", 38),
    "MA-20": ("Panne électrique ou capteur", "pannes", 32),
    "MA-30": ("Défaut automatisme", "pannes", 26),
    "CS-10": ("Changement de série", "cds", 28),
    "CS-20": ("Attente outillage", "cds", 22),
    "RG-10": ("Réglage qualité en cours de série", "reg", 12),
    "RG-20": ("Micro-arrêts, bourrage", "reg", 9),
    "AP-10": ("Rupture d'approvisionnement", "mat", 24),
    "AP-20": ("Composant non conforme", "mat", 20),
    "MO-10": ("Absence non remplacée", "per", 26),
    "MO-20": ("Accompagnement, formation au poste", "per", 18),
    "LO-10": ("Attente cariste, évacuation en-cours", "aut", 16),
    "LO-20": ("Saturation aval", "aut", 14),
}
SUBIS = [c for c, v in MOTIFS.items() if v[1] != "plan"]
CAD = {a[0]: a[3] for a in ARTICLES}
ART_LIGNE = {}
for a in ARTICLES:
    ART_LIGNE.setdefault(a[2], []).append(a[0])


def ouvres(a, b):
    out, d = [], a
    while d <= b:
        if d.weekday() < 5 and d not in FERIES and not (CONGES[0] <= d <= CONGES[1]):
            out.append(d)
        d += dt.timedelta(days=1)
    return out


def generer():
    prod, arrets = [], []
    jours = ouvres(DEB, FIN)
    for ji, jour in enumerate(jours):
        for ligne in LIGNES:
            for ei, eq in enumerate(EQUIPES[ligne]):
                k = ji * 7 + ei * 3 + LIGNES.index(ligne)
                art = ART_LIGNE[ligne][ji % len(ART_LIGNE[ligne])]
                cad = CAD[art]
                to = 480 if eq != "Nuit" else 450

                evts = [("PL-10", MOTIFS["PL-10"][2])]
                if ji % 6 == 2 and eq == "Matin":
                    evts.append(("PL-20", MOTIFS["PL-20"][2] + k % 25))
                if ji % 11 == 5 and eq == "Matin":
                    evts.append(("PL-30", 45))
                n_subis = 2 + (k * 3) % 4
                for s in range(n_subis):
                    code = SUBIS[(k * 5 + s * 3) % len(SUBIS)]
                    base = MOTIFS[code][2]
                    duree = max(4, int(base * (0.55 + random.random() * 0.95)))
                    if ji in (18, 42, 63) and code.startswith("MA"):
                        duree += 75                    # trois dérives marquées
                    evts.append((code, duree))

                t = dt.datetime.combine(jour, dt.time(6 if eq == "Matin" else
                                                      13 if eq == "Après-midi" else 21, 0))
                t += dt.timedelta(minutes=20)
                par_rub = {r: 0 for r in ("plan", "pannes", "cds", "reg", "mat", "per", "aut")}
                npan = 0
                for code, duree in evts:
                    deb = t + dt.timedelta(minutes=random.randint(8, 40))
                    fin = deb + dt.timedelta(minutes=duree)
                    arrets.append({
                        "date": jour, "debut": deb, "fin": fin, "ligne": ligne,
                        "equipe": eq, "code": code, "duree": duree})
                    par_rub[MOTIFS[code][1]] += duree
                    if code.startswith("MA"):
                        npan += 1
                    t = fin

                planifie = par_rub["plan"]
                subi = sum(v for r, v in par_rub.items() if r != "plan")
                tr, tf = to - planifie, to - planifie - subi
                perf = 0.855 + ((k * 11) % 12) / 100.0
                produite = int(tf * cad * perf)
                rft = 0.978 + ((k * 7) % 18) / 1000.0
                conforme = int(produite * rft)
                retouche = int((produite - conforme) * 0.4)
                rebut = produite - conforme - retouche
                demande = int((conforme + retouche) / (0.93 + ((k * 5) % 8) / 100.0))
                eff = {"LA-01 Assemblage": 12, "LU-02 Usinage": 8,
                       "LC-03 Conditionnement": 6}[ligne]
                present = eff - (1 if (k % 9) == 0 else 0)
                prod.append({
                    "date": jour, "ligne": ligne, "equipe": eq, "art": art,
                    "des": next(a[1] for a in ARTICLES if a[0] == art),
                    "to": to, "plan": planifie, "pannes": par_rub["pannes"],
                    "cds": par_rub["cds"], "reg": par_rub["reg"], "mat": par_rub["mat"],
                    "per": par_rub["per"], "aut": par_rub["aut"], "npan": npan,
                    "cad": cad, "prod": produite, "conf": conforme, "ret": retouche,
                    "reb": rebut, "dem": demande, "effP": eff, "effR": present,
                    "acc": 1 if (jour == dt.date(2025, 11, 18) and ligne == LIGNES[0]
                                 and eq == "Matin") else 0,
                    "presq": 1 if (k % 13) == 0 else 0})
    return prod, arrets


def generer_ofs(ref):
    ofs, n = [], 0
    plan = [("Soldé", -110, -8, 96), ("En cours", -14, 5, 26),
            ("Planifié", 3, 58, 52), ("Retard", -26, -1, 14)]
    for statut, j0, j1, nb in plan:
        for i in range(nb):
            n += 1
            ligne = LIGNES[n % len(LIGNES)]
            art = ART_LIGNE[ligne][n % len(ART_LIGNE[ligne])]
            cad = CAD[art]
            deb = ref + dt.timedelta(days=j0 + int((j1 - j0) * i / max(1, nb - 1)))
            while deb.weekday() >= 5:
                deb += dt.timedelta(days=1)
            fin = deb + dt.timedelta(days=1 + n % 4)
            while fin.weekday() >= 5:
                fin += dt.timedelta(days=1)
            qte = int(cad * 60 * (4 + n % 10)) // 10 * 10
            debR = finR = None
            fait = 0
            if statut == "Soldé":
                debR = deb
                finR = fin + dt.timedelta(days=1 if n % 6 == 0 else 0)
                fait = qte - (n % 8) * 4
                st = "Soldé"
            elif statut == "En cours":
                debR = min(deb, ref)
                fait = int(qte * (0.1 + ((n * 19) % 75) / 100))
                st = "En cours"
            elif statut == "Retard":
                debR = deb
                fait = int(qte * (0.25 + ((n * 11) % 60) / 100))
                st = "En cours"
            else:
                st = "Planifié"
            ofs.append({
                "of": "OF%06d" % (250000 + n * 7), "ligne": ligne, "art": art,
                "des": next(a[1] for a in ARTICLES if a[0] == art),
                "client": CLIENTS[n % len(CLIENTS)], "qte": qte, "fait": max(0, fait),
                "debP": deb, "finP": fin, "debR": debR, "finR": finR, "statut": st,
                "prio": ["Normale", "Normale", "Normale", "Haute", "Urgente"][n % 5]})
    ofs.sort(key=lambda o: o["debP"])
    return ofs


# ---------------------------------------------------------------- écriture
FR = lambda d: d.strftime("%d/%m/%Y")
HM = lambda d: d.strftime("%H:%M")


def ecrire(nom, entetes, lignes):
    os.makedirs(OUT, exist_ok=True)
    chemin_csv = os.path.join(OUT, nom + ".csv")
    with open(chemin_csv, "w", encoding="utf-8-sig", newline="") as f:
        w = csv.writer(f, delimiter=";")
        w.writerow(entetes)
        w.writerows(lignes)
    wb = Workbook()
    ws = wb.active
    ws.title = nom[3:][:28]
    ws.append(entetes)
    for l in lignes:
        ws.append(l)
    for c in ws[1]:
        c.font = Font(name="Calibri", size=11, bold=True, color="FFFFFF")
        c.fill = PatternFill("solid", fgColor="44546A")
        c.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
    ws.row_dimensions[1].height = 30
    for i, h in enumerate(entetes, 1):
        ws.column_dimensions[gcl(i)].width = min(26, max(11, len(h) + 2))
    ws.freeze_panes = "A2"
    ws.auto_filter.ref = ws.dimensions
    wb.save(os.path.join(OUT, nom + ".xlsx"))
    return len(lignes)


def main():
    prod, arrets = generer()
    ref = max(p["date"] for p in prod)
    ofs = generer_ofs(ref)

    n1 = ecrire("01_production_journaliere",
        ["Date de production", "Centre de charge", "Equipe", "Code article",
         "Désignation article", "Temps d'ouverture (mn)", "Arrêts planifiés (mn)",
         "Arrêts panne (mn)", "Arrêts changement de série (mn)", "Arrêts réglage (mn)",
         "Arrêts manque composant (mn)", "Arrêts effectif (mn)", "Autres arrêts (mn)",
         "Nb interventions maintenance", "Cadence théorique (p/mn)", "Qté lancée",
         "Qté bonne 1er passage", "Qté retouchée", "Qté rebutée", "Qté au programme",
         "Effectif théorique", "Effectif présent", "Accidents AT", "Presqu'accidents"],
        [[FR(p["date"]), p["ligne"], p["equipe"], p["art"], p["des"], p["to"], p["plan"],
          p["pannes"], p["cds"], p["reg"], p["mat"], p["per"], p["aut"], p["npan"],
          str(p["cad"]).replace(".", ","), p["prod"], p["conf"], p["ret"], p["reb"],
          p["dem"], p["effP"], p["effR"], p["acc"], p["presq"]] for p in prod])

    n2 = ecrire("02_journal_arrets",
        ["Date", "Heure début", "Heure fin", "Centre de charge", "Equipe",
         "Code motif d'arrêt", "Libellé motif", "Durée (mn)"],
        [[FR(a["date"]), HM(a["debut"]), HM(a["fin"]), a["ligne"], a["equipe"],
          a["code"], MOTIFS[a["code"]][0], a["duree"]] for a in arrets])

    n3 = ecrire("03_ordres_fabrication",
        ["N° OF", "Centre de charge", "Code article", "Désignation", "Client",
         "Qté à lancer", "Qté déclarée", "Date début prévue", "Date fin prévue",
         "Date début réelle", "Date fin réelle", "Statut OF", "Priorité"],
        [[o["of"], o["ligne"], o["art"], o["des"], o["client"], o["qte"], o["fait"],
          FR(o["debP"]), FR(o["finP"]), FR(o["debR"]) if o["debR"] else "",
          FR(o["finR"]) if o["finR"] else "", o["statut"], o["prio"]] for o in ofs])

    print("site :", SITE, "| période :", FR(DEB), "→", FR(FIN), "| référence :", FR(ref))
    print("  01_production_journaliere : %4d lignes" % n1)
    print("  02_journal_arrets         : %4d lignes" % n2)
    print("  03_ordres_fabrication     : %4d lignes" % n3)
    return prod, arrets, ofs, ref


if __name__ == "__main__":
    main()
