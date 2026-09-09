# -*- coding: utf-8 -*-
"""Onglet CALC : moteur de calcul du cockpit (series, Pareto, cascade, historique)."""
from openpyxl.styles import Alignment
from common import *
import data as D

SP = "SAISIE_PROD"
SR, SL = 5, 204                    # plage de donnees SAISIE_PROD
AR_F, AR_L = 5, 404                # plage de donnees ARRETS

JOUR0, JOURN = 12, 42              # serie journaliere
PAR0, PARN = 46, 53                # Pareto (8 rubriques)
CAS0, CASN = 57, 69                # cascade des pertes
LIG0, LIGN = 72, 76                # TRS par ligne
EQU0, EQUN = 79, 83                # TRS par equipe
MOI0, MOIN = 88, 100               # historique 12 mois (+ mois courant)
ACT0, ACTN = 103, 107              # repartition du plan d'actions
CAU0, CAUN = 110, 149              # causes d'arret (referentiel)
TOP0, TOPN = 152, 156              # top 5 des causes
GAU0, GAUN = 161, 164              # jauges du management visuel
ACH0, ACHN = 168, 287              # score des actions ouvertes
TR3, TR3N = 290, 292               # top 3 des actions

TOT = {                            # ligne du bloc de totaux -> colonne SAISIE_PROD
    8:  ("Temps d'ouverture TO (min)",        "F"),
    9:  ("Arrêts planifiés (min)",            "G"),
    10: ("Pannes (min)",                      "H"),
    11: ("Changement de série (min)",         "I"),
    12: ("Réglages et micro-arrêts (min)",    "J"),
    13: ("Manque matière (min)",              "K"),
    14: ("Manque personnel (min)",            "L"),
    15: ("Autres arrêts subis (min)",         "M"),
    16: ("Perte de cadence (min)",            "AP"),
    17: ("Pertes non-qualité (min)",          "AQ"),
    18: ("Temps requis TR (min)",             "AB"),
    19: ("Temps de marche TF (min)",          "AD"),
    20: ("Temps utile TU (min)",              "AE"),
    22: ("Quantité produite",                 "P"),
    23: ("Quantité conforme 1er passage",     "Q"),
    24: ("Quantité retouchée",                "R"),
    25: ("Quantité demandée",                 "T"),
    26: ("Accidents avec arrêt",              "X"),
    27: ("Nombre de pannes",                  "N"),
    28: ("Effectif prévu (cumul)",            "U"),
    29: ("Effectif présent (cumul)",          "V"),
}

RUBRIQUES = [
    ("Pannes machine",              "H"),
    ("Changement de série",         "I"),
    ("Réglages et micro-arrêts",    "J"),
    ("Manque matière",              "K"),
    ("Manque personnel",            "L"),
    ("Autres arrêts subis",         "M"),
    ("Perte de cadence",            "AP"),
    ("Pertes non-qualité",          "AQ"),
]

CASCADE = [
    ("Temps d'ouverture",          "0",                      "$AC$8",  1),
    ("Arrêts planifiés",           "$AC$8-$AC$9",              "$AC$9",  0),
    ("Temps requis",               "0",                      "$AC$18", 1),
    ("Pannes machine",             "$AC$18-$AC$10",            "$AC$10", 0),
    ("Changement de série",        "$AC$18-SUM($AC$10:$AC$11)", "$AC$11", 0),
    ("Réglages, micro-arrêts",     "$AC$18-SUM($AC$10:$AC$12)", "$AC$12", 0),
    ("Manque matière",             "$AC$18-SUM($AC$10:$AC$13)", "$AC$13", 0),
    ("Manque personnel",           "$AC$18-SUM($AC$10:$AC$14)", "$AC$14", 0),
    ("Autres arrêts subis",        "$AC$18-SUM($AC$10:$AC$15)", "$AC$15", 0),
    ("Temps de marche",            "0",                      "$AC$19", 1),
    ("Perte de cadence",           "$AC$19-$AC$16",            "$AC$16", 0),
    ("Pertes non-qualité",         "$AC$19-SUM($AC$16:$AC$17)", "$AC$17", 0),
    ("Temps utile",                "0",                      "$AC$20", 1),
]


def _sumifs(col, crit_date):
    """SUMIFS sur SAISIE_PROD avec filtre cockpit."""
    return ("SUMIFS({s}!${c}${f}:${c}${l},{s}!$A${f}:$A${l},{d},"
            "{s}!$AS${f}:$AS${l},1)").format(s=SP, c=col, f=SR, l=SL, d=crit_date)


def _sumifs_periode(col):
    return ('SUMIFS({s}!${c}${f}:${c}${l},{s}!$A${f}:$A${l},">="&$B$7,'
            '{s}!$A${f}:$A${l},"<="&$B$8,{s}!$AS${f}:$AS${l},1)').format(
        s=SP, c=col, f=SR, l=SL)


def build(wb):
    ws = wb.create_sheet("CALC")
    ws.sheet_view.showGridLines = False
    ws.sheet_properties.tabColor = "8A9AA8"

    title_band(ws, 1, 20, "MOTEUR DE CALCUL  —  feuille technique, ne pas modifier",
               "Cette feuille agrège les données de saisie selon les filtres du cockpit et alimente l'ensemble des graphiques. "
               "Toute modification manuelle fausse le tableau de bord.")

    def h(ref, txt, size=9.5):
        label(ws, ref, txt, bold=True, size=size, color=STEEL)

    def cell(ref, value, fmt=None, bold=False, color="1A1A1A"):
        c = ws[ref]
        c.value = value
        c.font = f(9, bold, color)
        if fmt:
            c.number_format = fmt
        c.alignment = Alignment(horizontal="center", vertical="center")
        return c

    # ---------------------------------------------------- filtres resolus
    h("A3", "PÉRIMÈTRE RÉSOLU")
    lignes = [("Année", "=COCKPIT!$B$4", "0"), ("Mois", "=COCKPIT!$D$4", "0"),
              ("Ligne", "=COCKPIT!$H$4", None), ("Équipe", "=COCKPIT!$K$4", None),
              ("Début de période", "=DATE($B$3,$B$4,1)", DATE),
              ("Fin de période", "=EOMONTH($B$7,0)", DATE),
              ("Jours de production saisis", "=COUNT($M$%d:$M$%d)" % (JOUR0, JOURN), "0")]
    for i, (lab, formule, fmt) in enumerate(lignes):
        r = 3 + i
        label(ws, "A%d" % r, lab, bold=False, size=9)
        cell("B%d" % r, formule, fmt)

    # ---------------------------------------------------- reperes statistiques
    h("AB2", "REPÈRES STATISTIQUES  (colonne AC)")
    reperes = [
        ("Cible TRS", '=INDEX(PARAMETRES!$D$16:$D$29,MATCH("TRS",PARAMETRES!$A$16:$A$29,0))', PCT),
        ("Moyenne des TRS journaliers", '=IFERROR(AVERAGE($M$%d:$M$%d),"")' % (JOUR0, JOURN), PCT),
        ("Écart-type des TRS journaliers", '=IFERROR(STDEV($M$%d:$M$%d),"")' % (JOUR0, JOURN), PCT),
        ("Limite de contrôle supérieure (moy. + 2σ)", '=IFERROR(MIN(1,$AC$4+2*$AC$5),"")', PCT),
        ("Limite de contrôle inférieure (moy. - 2σ)", '=IFERROR(MAX(0,$AC$4-2*$AC$5),"")', PCT),
    ]
    for i, (lab, formule, fmt) in enumerate(reperes):
        r = 3 + i
        label(ws, "AB%d" % r, lab, bold=False, size=9)
        cell("AC%d" % r, formule, fmt)

    # ---------------------------------------------------- totaux de periode
    for r, (lab, col) in TOT.items():
        label(ws, "AB%d" % r, lab, bold=False, size=9)
        cell("AC%d" % r, "=" + _sumifs_periode(col), NUM)

    # ---------------------------------------------------- seuils SQCDP
    h("AB31", "SEUILS UTILISÉS PAR LE MANAGEMENT VISUEL")
    seuils = [("QUAL", "Cible qualité RFT", "D"), ("QUAL", "Alerte qualité RFT", "E"),
              ("TRS", "Cible TRS", "D"), ("TRS", "Alerte TRS", "E"),
              ("SERVICE", "Cible taux de service", "D"), ("SERVICE", "Alerte taux de service", "E"),
              ("PRESENCE", "Cible taux de présence", "D"), ("PRESENCE", "Alerte taux de présence", "E")]
    for i, (code, lab, col) in enumerate(seuils):
        r = 32 + i
        label(ws, "AB%d" % r, lab, bold=False, size=9)
        cell("AC%d" % r,
             '=INDEX(PARAMETRES!${c}$16:${c}$29,MATCH("{k}",PARAMETRES!$A$16:$A$29,0))'.format(c=col, k=code),
             PCT)

    # ---------------------------------------------------- serie journaliere
    band(ws, JOUR0 - 2, 1, 27, "SÉRIE JOURNALIÈRE DU MOIS SÉLECTIONNÉ  —  valeurs, séries de graphique et statuts du management visuel")
    entetes = ["Jour", "Date", "Temps utile", "Temps requis", "Temps de marche",
               "Temps d'ouverture", "Qté produite", "Qté conforme", "Qté retouchée",
               "Qté demandée", "Accidents", "Arrêts subis", "TRS", "Cible",
               "Moyenne mobile 7 j", "LCS", "LCI", "TRS (graphique)",
               "Moyenne mobile (graphique)", "Effectif prévu", "Effectif présent",
               "Presqu'accidents", "Statut S", "Statut Q", "Statut C", "Statut D", "Statut P"]
    header_row(ws, JOUR0 - 1, entetes, height=34)
    for i in range(JOURN - JOUR0 + 1):
        r = JOUR0 + i
        d = "$B%d" % r
        ws["A%d" % r] = i + 1
        ws["B%d" % r] = '=IF($A%d>DAY($B$8),"",DATE($B$3,$B$4,$A%d))' % (r, r)
        for col, src in (("C", "AE"), ("D", "AB"), ("E", "AD"), ("F", "F"), ("G", "P"),
                         ("H", "Q"), ("I", "R"), ("J", "T"), ("K", "X"), ("L", "AC"),
                         ("T", "U"), ("U", "V"), ("V", "Y")):
            ws["%s%d" % (col, r)] = '=IF(%s="","",%s)' % (d, _sumifs(src, d))
        ws["M%d" % r] = '=IF(OR($B%d="",$D%d=0),"",$C%d/$D%d)' % (r, r, r, r)
        ws["N%d" % r] = '=IF($B%d="",NA(),$AC$3)' % r
        s = max(JOUR0, r - 6)
        ws["O%d" % r] = '=IF(COUNT($M%d:$M%d)<3,"",AVERAGE($M%d:$M%d))' % (s, r, s, r)
        ws["P%d" % r] = '=IF(OR($B%d="",$AC$5=""),NA(),$AC$6)' % r
        ws["Q%d" % r] = '=IF(OR($B%d="",$AC$5=""),NA(),$AC$7)' % r
        ws["R%d" % r] = '=IF($M%d="",NA(),$M%d)' % (r, r)
        ws["S%d" % r] = '=IF($O%d="",NA(),$O%d)' % (r, r)
        # statuts du management visuel : 1 vert, 2 orange, 3 rouge, "" sans production
        prod = '$D%d=0' % r
        ws["W%d" % r] = ('=IF(OR($B{r}="",{p}),"",IF($K{r}>0,3,IF($V{r}>0,2,1)))'
                         .format(r=r, p=prod))
        ws["X%d" % r] = ('=IF(OR($B{r}="",{p},$G{r}=0),"",IF($H{r}/$G{r}>=$AC$32,1,'
                         'IF($H{r}/$G{r}>=$AC$33,2,3)))').format(r=r, p=prod)
        ws["Y%d" % r] = ('=IF(OR($B{r}="",{p}),"",IF($M{r}>=$AC$34,1,IF($M{r}>=$AC$35,2,3)))'
                         .format(r=r, p=prod))
        ws["Z%d" % r] = ('=IF(OR($B{r}="",{p},$J{r}=0),"",IF(($H{r}+$I{r})/$J{r}>=$AC$36,1,'
                         'IF(($H{r}+$I{r})/$J{r}>=$AC$37,2,3)))').format(r=r, p=prod)
        ws["AA%d" % r] = ('=IF(OR($B{r}="",{p},$T{r}=0),"",IF($U{r}/$T{r}>=$AC$38,1,'
                          'IF($U{r}/$T{r}>=$AC$39,2,3)))').format(r=r, p=prod)
        for col in ("A B C D E F G H I J K L M N O P Q R S T U V W X Y Z AA").split():
            c = ws["%s%d" % (col, r)]
            c.font = f(9)
            c.alignment = Alignment(horizontal="center", vertical="center")
            c.border = BOX
            if col == "B":
                c.number_format = DATE
            elif col in ("M", "N", "O", "P", "Q", "R", "S"):
                c.number_format = PCT
            elif col in ("W", "X", "Y", "Z", "AA"):
                c.number_format = "0"
            elif col != "A":
                c.number_format = NUM
        ws.row_dimensions[r].height = 14

    # ---------------------------------------------------- Pareto
    band(ws, PAR0 - 2, 1, 12, "PARETO DES PERTES DE TEMPS SUBIES  (hors arrêts planifiés)")
    header_row(ws, PAR0 - 1, ["Rubrique de perte", "Minutes perdues", "Valeur départagée",
                              "Rang", "", "Rubrique classée", "Minutes", "Part",
                              "Part cumulée", "", "", ""], height=28)
    tot_par = "SUM($B$%d:$B$%d)" % (PAR0, PARN)
    for i, (lib, col) in enumerate(RUBRIQUES):
        r = PAR0 + i
        ws["A%d" % r] = lib
        ws["B%d" % r] = "=" + _sumifs_periode(col)
        ws["C%d" % r] = '=$B%d+(%d-ROW())/100000' % (r, PARN + 1)
        ws["D%d" % r] = '=RANK($C%d,$C$%d:$C$%d)' % (r, PAR0, PARN)
        k = i + 1
        ws["F%d" % r] = ('=INDEX($A$%d:$A$%d,MATCH(LARGE($C$%d:$C$%d,%d),$C$%d:$C$%d,0))'
                         % (PAR0, PARN, PAR0, PARN, k, PAR0, PARN))
        ws["G%d" % r] = ('=INDEX($B$%d:$B$%d,MATCH(LARGE($C$%d:$C$%d,%d),$C$%d:$C$%d,0))'
                         % (PAR0, PARN, PAR0, PARN, k, PAR0, PARN))
        ws["H%d" % r] = '=IFERROR($G%d/%s,"")' % (r, tot_par)
        ws["I%d" % r] = '=IFERROR(SUM($G$%d:$G%d)/%s,"")' % (PAR0, r, tot_par)
        for col2, fmt in (("A", None), ("B", NUM), ("C", "0.00000"), ("D", "0"),
                          ("F", None), ("G", NUM), ("H", PCT), ("I", PCT)):
            c = ws["%s%d" % (col2, r)]
            c.font = f(9)
            c.border = BOX
            c.number_format = fmt or "General"
            c.alignment = Alignment(horizontal="left" if fmt is None else "center",
                                    vertical="center", indent=1)
        ws.row_dimensions[r].height = 14

    # ---------------------------------------------------- cascade des pertes
    band(ws, CAS0 - 2, 1, 12, "CASCADE DES PERTES DE TEMPS  (du temps d'ouverture au temps utile)")
    header_row(ws, CAS0 - 1, ["Étape", "Socle (invisible)", "Valeur affichée",
                              "Jalon", "", "", "", "", "", "", "", ""], height=26)
    for i, (lib, base, val, jalon) in enumerate(CASCADE):
        r = CAS0 + i
        ws["A%d" % r] = lib
        ws["B%d" % r] = "=" + base if base != "0" else 0
        ws["C%d" % r] = "=" + val
        ws["D%d" % r] = jalon
        for col2 in "ABCD":
            c = ws["%s%d" % (col2, r)]
            c.font = f(9, bool(jalon))
            c.border = BOX
            c.number_format = NUM if col2 in "BC" else "General"
            c.alignment = Alignment(horizontal="left" if col2 == "A" else "center",
                                    vertical="center", indent=1)
        ws.row_dimensions[r].height = 14

    # ---------------------------------------------------- TRS par ligne / equipe
    def bloc_axe(r0, rn, titre, ref_col, saisie_col, param_col):
        band(ws, r0 - 2, 1, 12, titre)
        header_row(ws, r0 - 1, ["Libellé", "Temps utile", "Temps requis", "TRS",
                                "TRS (graphique)", "", "", "", "", "", "", ""], height=26)
        for i in range(rn - r0 + 1):
            r = r0 + i
            ws["A%d" % r] = ('=IF(PARAMETRES!${p}{pr}="","",PARAMETRES!${p}{pr})'
                             .format(p=param_col, pr=7 + i))
            for col, src in (("B", "AE"), ("C", "AB")):
                ws["%s%d" % (col, r)] = (
                    '=IF($A{r}="","",SUMIFS({s}!${c}${f}:${c}${l},{s}!$A${f}:$A${l},">="&$B$7,'
                    '{s}!$A${f}:$A${l},"<="&$B$8,{s}!${x}${f}:${x}${l},$A{r}))'
                ).format(r=r, s=SP, c=src, f=SR, l=SL, x=saisie_col)
            ws["D%d" % r] = '=IF(OR($A{r}="",$C{r}=0),"",$B{r}/$C{r})'.format(r=r)
            ws["E%d" % r] = '=IF($D{r}="",NA(),$D{r})'.format(r=r)
            for col2, fmt in (("A", None), ("B", NUM), ("C", NUM), ("D", PCT), ("E", PCT)):
                c = ws["%s%d" % (col2, r)]
                c.font = f(9)
                c.border = BOX
                c.number_format = fmt or "General"
                c.alignment = Alignment(horizontal="left" if fmt is None else "center",
                                        vertical="center", indent=1)
            ws.row_dimensions[r].height = 14

    bloc_axe(LIG0, LIGN, "TRS PAR LIGNE DE PRODUCTION  (mois sélectionné — comparaison entre lignes, hors filtres du cockpit)", "I", "C", "I")
    bloc_axe(EQU0, EQUN, "TRS PAR ÉQUIPE  (mois sélectionné — comparaison entre équipes, hors filtres du cockpit)", "J", "B", "J")

    # ---------------------------------------------------- historique mensuel
    band(ws, MOI0 - 2, 1, 19, "HISTORIQUE SUR 13 MOIS GLISSANTS  (même périmètre de filtrage)")
    header_row(ws, MOI0 - 1, ["Mois", "Date de début", "Temps utile", "Temps requis",
                              "Temps de marche", "Temps d'ouverture", "Qté produite",
                              "Qté conforme", "Qté retouchée", "Qté demandée", "Accidents",
                              "TRS", "Disponibilité", "Performance", "Qualité RFT", "TRG",
                              "Taux de service", "TRS (graphique)", "Cible (graphique)"],
               height=32)
    for i in range(MOIN - MOI0 + 1):
        r = MOI0 + i
        k = r - MOIN
        ws["B%d" % r] = '=EDATE($B$7,%d)' % k
        ws["A%d" % r] = ('=INDEX(PARAMETRES!$AM$7:$AM$18,MONTH($B{r}))&" "&TEXT(YEAR($B{r}),"0000")'
                         .format(r=r))
        for col, src in (("C", "AE"), ("D", "AB"), ("E", "AD"), ("F", "F"), ("G", "P"),
                         ("H", "Q"), ("I", "R"), ("J", "T"), ("K", "X")):
            ws["%s%d" % (col, r)] = (
                '=SUMIFS({s}!${c}${f}:${c}${l},{s}!$A${f}:$A${l},">="&$B{r},'
                '{s}!$A${f}:$A${l},"<"&EDATE($B{r},1),{s}!$AS${f}:$AS${l},1)'
            ).format(s=SP, c=src, f=SR, l=SL, r=r)
        ws["L%d" % r] = '=IFERROR($C{r}/$D{r},"")'.format(r=r)
        ws["M%d" % r] = '=IFERROR($E{r}/$D{r},"")'.format(r=r)
        ws["N%d" % r] = '=IFERROR($C{r}*($G{r}/$H{r})/$E{r},"")'.format(r=r)
        ws["O%d" % r] = '=IFERROR($H{r}/$G{r},"")'.format(r=r)
        ws["P%d" % r] = '=IFERROR($C{r}/$F{r},"")'.format(r=r)
        ws["Q%d" % r] = '=IFERROR(($H{r}+$I{r})/$J{r},"")'.format(r=r)
        ws["R%d" % r] = '=IF($L{r}="",NA(),$L{r})'.format(r=r)
        ws["S%d" % r] = '=$AC$3'
        for col2 in "ABCDEFGHIJKLMNOPQRS":
            c = ws["%s%d" % (col2, r)]
            c.font = f(9, k == 0)
            c.border = BOX
            c.alignment = Alignment(horizontal="center", vertical="center")
            if col2 == "B":
                c.number_format = "mmm yyyy"
            elif col2 in "LMNOPQRS":
                c.number_format = PCT
            elif col2 != "A":
                c.number_format = NUM
        ws["A%d" % r].alignment = Alignment(horizontal="left", vertical="center", indent=1)
        ws.row_dimensions[r].height = 14

    # ---------------------------------------------------- plan d'actions
    band(ws, ACT0 - 2, 1, 12, "RÉPARTITION DU PLAN D'ACTIONS PAR STATUT")
    header_row(ws, ACT0 - 1, ["Statut", "Nombre d'actions", "", "", "", "", "", "", "", "", "", ""],
               height=24)
    for i, st in enumerate(D.STATUTS):
        r = ACT0 + i
        ws["A%d" % r] = st
        ws["B%d" % r] = '=COUNTIF(PLAN_ACTIONS!$Q$9:$Q$128,$A%d)' % r
        for col2 in "AB":
            c = ws["%s%d" % (col2, r)]
            c.font = f(9)
            c.border = BOX
            c.alignment = Alignment(horizontal="left" if col2 == "A" else "center",
                                    vertical="center", indent=1)
        ws.row_dimensions[r].height = 14

    # ---------------------------------------------------- causes d'arret
    band(ws, CAU0 - 2, 1, 12, "ANALYSE DES CAUSES D'ARRÊT SUBI  (source : onglet ARRETS, hors arrêts planifiés)")
    header_row(ws, CAU0 - 1, ["Code", "Libellé", "Minutes", "Occurrences", "Durée moyenne",
                              "Valeur départagée", "", "", "", "", "", ""], height=26)
    for i in range(CAUN - CAU0 + 1):
        r = CAU0 + i
        ws["A%d" % r] = '=IF(PARAMETRES!$R%d="","",PARAMETRES!$R%d)' % (7 + i, 7 + i)
        ws["B%d" % r] = '=IF($A{r}="","",PARAMETRES!$S{p})'.format(r=r, p=7 + i)
        ws["C%d" % r] = ('=IF($A{r}="",0,SUMIFS(ARRETS!$G${f}:$G${l},ARRETS!$A${f}:$A${l},">="&$B$7,'
                         'ARRETS!$A${f}:$A${l},"<="&$B$8,ARRETS!$H${f}:$H${l},$A{r},'
                         'ARRETS!$Q${f}:$Q${l},1,ARRETS!$K${f}:$K${l},"<>Planifié"))').format(r=r, f=AR_F, l=AR_L)
        ws["D%d" % r] = ('=IF($A{r}="",0,COUNTIFS(ARRETS!$A${f}:$A${l},">="&$B$7,'
                         'ARRETS!$A${f}:$A${l},"<="&$B$8,ARRETS!$H${f}:$H${l},$A{r},'
                         'ARRETS!$Q${f}:$Q${l},1,ARRETS!$K${f}:$K${l},"<>Planifié"))').format(r=r, f=AR_F, l=AR_L)
        ws["E%d" % r] = '=IFERROR($C{r}/$D{r},"")'.format(r=r)
        ws["F%d" % r] = '=$C{r}+({n}-ROW())/100000'.format(r=r, n=CAUN + 1)
        for col2, fmt in (("A", None), ("B", None), ("C", NUM), ("D", NUM), ("E", NUM1),
                          ("F", "0.00000")):
            c = ws["%s%d" % (col2, r)]
            c.font = f(9)
            c.border = BOX
            c.number_format = fmt or "General"
            c.alignment = Alignment(horizontal="left" if fmt is None else "center",
                                    vertical="center", indent=1)
        ws.row_dimensions[r].height = 14

    band(ws, TOP0 - 2, 1, 12, "TOP 5 DES CAUSES D'ARRÊT SUBI DE LA PÉRIODE")
    header_row(ws, TOP0 - 1, ["Rang", "Code", "Libellé de la cause", "Minutes perdues",
                              "Occurrences", "Durée moyenne", "Part des arrêts", "",
                              "", "", "", ""], height=26)
    for i in range(TOPN - TOP0 + 1):
        r = TOP0 + i
        k = i + 1
        m = 'MATCH(LARGE($F$%d:$F$%d,%d),$F$%d:$F$%d,0)' % (CAU0, CAUN, k, CAU0, CAUN)
        ws["A%d" % r] = k
        ws["B%d" % r] = '=IFERROR(INDEX($A$%d:$A$%d,%s),"")' % (CAU0, CAUN, m)
        ws["C%d" % r] = '=IFERROR(INDEX($B$%d:$B$%d,%s),"")' % (CAU0, CAUN, m)
        ws["D%d" % r] = '=IFERROR(INDEX($C$%d:$C$%d,%s),"")' % (CAU0, CAUN, m)
        ws["E%d" % r] = '=IFERROR(INDEX($D$%d:$D$%d,%s),"")' % (CAU0, CAUN, m)
        ws["F%d" % r] = '=IFERROR($D%d/$E%d,"")' % (r, r)
        ws["G%d" % r] = '=IFERROR($D%d/SUM($C$%d:$C$%d),"")' % (r, CAU0, CAUN)
        for col2, fmt in (("A", "0"), ("B", None), ("C", None), ("D", NUM), ("E", NUM),
                          ("F", NUM1), ("G", PCT)):
            c = ws["%s%d" % (col2, r)]
            c.font = f(9)
            c.border = BOX
            c.number_format = fmt or "General"
            c.alignment = Alignment(horizontal="left" if fmt is None else "center",
                                    vertical="center", indent=1)
        ws.row_dimensions[r].height = 14

    # ---------------------------------------------------- jauges
    band(ws, GAU0 - 2, 1, 12, "JAUGES DU MANAGEMENT VISUEL  (valeur, reste, demi-cercle masqué)")
    header_row(ws, GAU0 - 1, ["Indicateur", "Valeur", "Reste", "Masqué", "Cible",
                              "Libellé affiché", "", "", "", "", "", ""], height=24)
    jauges = [("Taux de rendement synthétique", '=IFERROR($AC$20/$AC$18,0)', "TRS"),
              ("Qualité au premier passage", '=IFERROR($AC$23/$AC$22,0)', "QUAL"),
              ("Taux de service", '=IFERROR(($AC$23+$AC$24)/$AC$25,0)', "SERVICE"),
              ("Taux de présence", '=IFERROR($AC$29/$AC$28,0)', "PRESENCE")]
    for i, (lib, formule, code) in enumerate(jauges):
        r = GAU0 + i
        ws["A%d" % r] = lib
        ws["B%d" % r] = formule
        ws["C%d" % r] = '=MAX(0,1-$B%d)' % r
        ws["D%d" % r] = 1
        ws["E%d" % r] = ('=INDEX(PARAMETRES!$D$16:$D$29,MATCH("%s",PARAMETRES!$A$16:$A$29,0))'
                         % code)
        ws["F%d" % r] = '=TEXT($B{r},"0.0%")&"   |   cible "&TEXT($E{r},"0.0%")'.format(r=r)
        for col2, fmt in (("A", None), ("B", PCT), ("C", PCT), ("D", "0"), ("E", PCT), ("F", None)):
            c = ws["%s%d" % (col2, r)]
            c.font = f(9)
            c.border = BOX
            c.number_format = fmt or "General"
            c.alignment = Alignment(horizontal="left" if fmt is None else "center",
                                    vertical="center", indent=1)
        ws.row_dimensions[r].height = 14

    # ---------------------------------------------------- priorisation des actions
    band(ws, ACH0 - 2, 1, 12, "PRIORISATION DES ACTIONS OUVERTES  (source : PLAN_ACTIONS)")
    header_row(ws, ACH0 - 1, ["Criticité si action ouverte", "Valeur départagée", "", "", "",
                              "", "", "", "", "", "", ""], height=24)
    for i in range(ACHN - ACH0 + 1):
        r = ACH0 + i
        pa = 9 + i
        ws["A%d" % r] = ('=IF(PLAN_ACTIONS!$E{p}="",0,IF(OR(PLAN_ACTIONS!$Q{p}="Terminée",'
                         'PLAN_ACTIONS!$Q{p}="Abandonnée"),0,IF(PLAN_ACTIONS!$I{p}="",0,'
                         'PLAN_ACTIONS!$I{p})))').format(p=pa)
        ws["B%d" % r] = '=$A{r}+({n}-ROW())/100000'.format(r=r, n=ACHN + 1)
        for col2 in "AB":
            c = ws["%s%d" % (col2, r)]
            c.font = f(9)
            c.border = BOX
            c.number_format = "0.00000"
            c.alignment = Alignment(horizontal="center", vertical="center")
        ws.row_dimensions[r].height = 13

    band(ws, TR3 - 2, 1, 12, "TROIS ACTIONS PRIORITAIRES  (criticité la plus élevée parmi les actions ouvertes)")
    header_row(ws, TR3 - 1, ["Rang", "N°", "Problème constaté", "Pilote", "Date cible",
                             "Avancement", "Criticité", "Retard (j)", "Action décidée",
                             "", "", ""], height=24)
    for i in range(TR3N - TR3 + 1):
        r = TR3 + i
        m = 'MATCH(LARGE($B$%d:$B$%d,%d),$B$%d:$B$%d,0)' % (ACH0, ACHN, i + 1, ACH0, ACHN)
        ok = 'LARGE($A$%d:$A$%d,%d)=0' % (ACH0, ACHN, i + 1)
        srcs = [("B", "A"), ("C", "E"), ("D", "N"), ("E", "O"), ("F", "P"), ("G", "I"), ("H", "S"),
                ("I", "L")]
        ws["A%d" % r] = i + 1
        for dst, src in srcs:
            ws["%s%d" % (dst, r)] = ('=IF({ok},"",IFERROR(INDEX(PLAN_ACTIONS!${s}$9:${s}$128,{m}),""))'
                                     .format(ok=ok, s=src, m=m))
        for col2, fmt in (("A", "0"), ("B", None), ("C", None), ("D", None), ("E", DATE),
                          ("F", PCT), ("G", "0"), ("H", "0"), ("I", None)):
            c = ws["%s%d" % (col2, r)]
            c.font = f(9)
            c.border = BOX
            c.number_format = fmt or "General"
            c.alignment = Alignment(horizontal="left" if fmt is None else "center",
                                    vertical="center", indent=1)
        ws.row_dimensions[r].height = 14

    widths(ws, {"A": 30, "B": 14, "C": 13, "D": 13, "E": 14, "F": 26, "G": 13, "H": 12,
                "I": 13, "J": 13, "K": 11, "L": 13, "M": 11, "N": 11, "O": 15, "P": 10,
                "Q": 10, "R": 14, "S": 16, "T": 12, "U": 12, "V": 13,
                "W": 9, "X": 9, "Y": 9, "Z": 9, "AA": 9, "AB": 34, "AC": 14})
    page(ws)
    return ws
