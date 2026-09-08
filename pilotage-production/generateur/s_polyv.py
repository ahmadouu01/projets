# -*- coding: utf-8 -*-
"""Onglet POLYVALENCE : matrice de competences ILUO et couverture des postes."""
from openpyxl.styles import Alignment, Border, Side
from openpyxl.worksheet.datavalidation import DataValidation
from openpyxl.formatting.rule import FormulaRule
from openpyxl.chart import BarChart, Reference
from openpyxl.chart.series import SeriesLabel
from openpyxl.chart.shapes import GraphicalProperties
from openpyxl.utils import get_column_letter as gcl
from common import *
import data as D

P1, P2 = 5, 16          # colonnes E..P : 12 postes
FIRST, LAST = 12, 51    # collaborateurs
SYN = 53                # premiere ligne de synthese


def build(wb):
    ws = wb.create_sheet("POLYVALENCE")
    ws.sheet_view.showGridLines = False
    ws.sheet_properties.tabColor = "2E6B4F"
    NC = 20

    title_band(ws, 1, NC, "MATRICE DE POLYVALENCE  —  couverture des postes de travail",
               "Échelle ILUO :  0 non formé  •  1 en formation  •  2 autonome avec appui  •  3 autonome  •  4 autonome et formateur. "
               "Un poste est considéré couvert lorsque l'effectif autonome (niveau ≥ 3) atteint la cible définie dans PARAMETRES.")

    tuiles = [
        ("Effectif suivi", 1, 3, '=COUNTA($B${f}:$B${l})'.format(f=FIRST, l=LAST), NUM),
        ("Taux de polyvalence global", 4, 6,
         '=IFERROR(SUM($E${f}:$P${l})/(4*COUNTIF($E$9:$P$9,"?*")*COUNTA($B${f}:$B${l})),"")'.format(f=FIRST, l=LAST), PCT),
        ("Postes sous la cible", 7, 9, '=COUNTIF($E${s}:$P${s},"<0")'.format(s=SYN + 3), NUM),
        ("Postes en mono-compétence", 10, 12, '=COUNTIF($E${s}:$P${s},"OUI")'.format(s=SYN + 5), NUM),
        ("Niveau moyen de maîtrise", 13, 16,
         '=IFERROR(AVERAGE($E${f}:$P${l}),"")'.format(f=FIRST, l=LAST), "0.00"),
        ("Postes critiques couverts", 17, 20,
         '=SUMPRODUCT(($E$10:$P$10=3)*($E${s}:$P${s}>=$E$11:$P$11))'.format(s=SYN), NUM),
    ]
    band(ws, 4, 1, NC, "INDICATEURS DE COMPÉTENCE  —  calculés automatiquement")
    for lib, c1, c2, formule, fmt in tuiles:
        ws.merge_cells(start_row=5, start_column=c1, end_row=5, end_column=c2)
        c = ws.cell(row=5, column=c1, value=lib)
        c.font = f(8.5, True, WHITE)
        c.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        for k in range(c1, c2 + 1):
            ws.cell(row=5, column=k).fill = fill(SLATE)
            ws.cell(row=5, column=k).border = BOX
        ws.merge_cells(start_row=6, start_column=c1, end_row=6, end_column=c2)
        v = ws.cell(row=6, column=c1, value=formule)
        v.font = f(15, True, NAVY)
        v.number_format = fmt
        v.alignment = Alignment(horizontal="center", vertical="center")
        for k in range(c1, c2 + 1):
            ws.cell(row=6, column=k).fill = fill(WHITE)
            ws.cell(row=6, column=k).border = BOX
    ws.row_dimensions[5].height = 26
    ws.row_dimensions[6].height = 26

    band(ws, 8, 1, NC, "MATRICE DES COMPÉTENCES")
    # lignes 9/10/11 : poste, criticite, cible (issus du referentiel)
    for lab, row, src in (("Poste de travail", 9, "X"), ("Criticité du poste (1 à 3)", 10, "Y"),
                          ("Effectif autonome cible", 11, "Z")):
        if row != 9:
            ws.merge_cells(start_row=row, start_column=1, end_row=row, end_column=4)
            c = ws.cell(row=row, column=1, value=lab)
            c.font = f(9, True, STEEL)
            c.alignment = Alignment(horizontal="right", vertical="center")
            for k in range(1, 5):
                ws.cell(row=row, column=k).fill = fill(LIGHT)
                ws.cell(row=row, column=k).border = BOX
            for k in range(17, 21):
                ws.cell(row=row, column=k).fill = fill(LIGHT)
                ws.cell(row=row, column=k).border = BOX
        for col in range(P1, P2 + 1):
            cc = ws.cell(row=row, column=col,
                         value=('=IF(INDEX(PARAMETRES!$X$7:$X$26,COLUMN()-{o})="","",'
                                'INDEX(PARAMETRES!${s}$7:${s}$26,COLUMN()-{o}))').format(
                             s=src, o=P1 - 1))
            cc.border = BOX
            cc.font = f(8.5, row == 9, WHITE if row == 9 else GREEN)
            cc.fill = fill(STEEL if row == 9 else LIGHT)
            cc.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
    ws.row_dimensions[9].height = 46
    ws.row_dimensions[10].height = 15
    ws.row_dimensions[11].height = 15

    entetes = ["Matricule", "Collaborateur", "Équipe", "Statut"] + [""] * 12 + \
              ["Postes\nautonomes", "Taux de\nmaîtrise", "Postes critiques\nmaîtrisés",
               "Prochaine étape de formation"]
    for i, h in enumerate(entetes):
        if i + 1 in range(P1, P2 + 1):
            continue
        c = ws.cell(row=9, column=i + 1, value=h)
        c.font = f(8.5, True, WHITE)
        c.fill = fill(STEEL)
        c.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        c.border = BOX

    for r in range(FIRST, LAST + 1):
        for col in range(1, 5):
            c = ws.cell(row=r, column=col)
            c.font = f(9, color="0000FF")
            c.fill = fill(INPUT_BG)
            c.border = BOX
            c.alignment = Alignment(horizontal="left", vertical="center", indent=1)
        for col in range(P1, P2 + 1):
            c = ws.cell(row=r, column=col)
            c.font = f(9, True, "0000FF")
            c.fill = fill(INPUT_BG)
            c.border = BOX
            c.alignment = Alignment(horizontal="center", vertical="center")
        ws["Q%d" % r] = '=IF($B{r}="","",COUNTIF($E{r}:$P{r},">=3"))'.format(r=r)
        ws["R%d" % r] = ('=IF($B{r}="","",IFERROR(SUM($E{r}:$P{r})/(4*COUNTIF($E$9:$P$9,"?*")),""))'
                         .format(r=r))
        ws["S%d" % r] = ('=IF($B{r}="","",SUMPRODUCT(($E$10:$P$10=3)*($E{r}:$P{r}>=3)))'
                         .format(r=r))
        for col, fmt in ((17, NUM), (18, PCT), (19, NUM)):
            c = ws.cell(row=r, column=col)
            c.font = f(9)
            c.fill = fill(CALC_BG)
            c.border = BOX
            c.number_format = fmt
            c.alignment = Alignment(horizontal="center", vertical="center")
        c = ws.cell(row=r, column=20)
        c.font = f(9, color="0000FF")
        c.fill = fill(INPUT_BG)
        c.border = BOX
        c.alignment = Alignment(horizontal="left", vertical="center", indent=1)
        ws.row_dimensions[r].height = 15

    for i, (mat, nom, eq, st) in enumerate(D.COLLABORATEURS):
        r = FIRST + i
        ws["A%d" % r], ws["B%d" % r], ws["C%d" % r], ws["D%d" % r] = mat, nom, eq, st
        for j, niv in enumerate(D.NIVEAUX[i]):
            ws.cell(row=r, column=P1 + j, value=niv)

    # ---------------------------------------------------------- synthese
    band(ws, SYN - 1, 1, NC, "COUVERTURE DES POSTES  —  lecture du risque de compétence")
    lignes = [
        ("Effectif autonome (niveau ≥ 3)", '=COUNTIF({c}${f}:{c}${l},">=3")', NUM),
        ("Effectif en formation (niveaux 1 et 2)", '=COUNTIFS({c}${f}:{c}${l},">=1",{c}${f}:{c}${l},"<3")', NUM),
        ("Effectif autonome cible", '={c}$11', NUM),
        ("Écart à la cible", '=IF({c}$11="","",{c}${s0}-{c}$11)', NUM),
        ("Taux de couverture", '=IFERROR({c}${s0}/{c}$11,"")', PCT),
        ("Risque de mono-compétence", '=IF({c}$9="","",IF({c}${s0}<=1,"OUI","non"))', None),
        ("Statut du poste", '=IF({c}$9="","",IF({c}${s0}>={c}$11,"Couvert",'
                            'IF({c}${s0}>={c}$11-1,"Vigilance","Critique")))', None),
    ]
    for i, (lab, formule, fmt) in enumerate(lignes):
        r = SYN + i
        ws.merge_cells(start_row=r, start_column=1, end_row=r, end_column=4)
        c = ws.cell(row=r, column=1, value=lab)
        c.font = f(9, True, STEEL)
        c.alignment = Alignment(horizontal="right", vertical="center")
        for k in range(1, 5):
            ws.cell(row=r, column=k).fill = fill(LIGHT)
            ws.cell(row=r, column=k).border = BOX
        for col in range(P1, P2 + 1):
            cl = gcl(col)
            cc = ws.cell(row=r, column=col,
                         value=formule.format(c="$" + cl, f=FIRST, l=LAST, s0=SYN))
            cc.font = f(9, True)
            cc.fill = fill(CALC_BG)
            cc.border = BOX
            cc.number_format = fmt if fmt else "General"
            cc.alignment = Alignment(horizontal="center", vertical="center")
        ws.row_dimensions[r].height = 16

    # validations et mises en forme conditionnelles
    dv = DataValidation(type="list", formula1='"0,1,2,3,4"', allow_blank=True,
                        showErrorMessage=True, errorTitle="Niveau ILUO",
                        error="Saisir un niveau de 0 à 4.")
    ws.add_data_validation(dv)
    dv.add("E%d:P%d" % (FIRST, LAST))
    dv2 = DataValidation(type="list", formula1="LST_EQUIPES", allow_blank=True)
    ws.add_data_validation(dv2)
    dv2.add("C%d:C%d" % (FIRST, LAST))

    couleurs = {0: ("FFFFFF", GREY), 1: ("FDE9A9", "7A5230"), 2: ("FFD8A8", "7A5230"),
                3: ("B7E4C0", GREEN), 4: ("6FBF8B", "0B3D1E")}
    for niv, (bg, fg) in couleurs.items():
        ws.conditional_formatting.add(
            "E%d:P%d" % (FIRST, LAST),
            FormulaRule(formula=['E%d=%d' % (FIRST, niv)], fill=fill(bg), font=f(9, True, fg)))
    ws.conditional_formatting.add(
        "E%d:P%d" % (SYN + 6, SYN + 6),
        FormulaRule(formula=['E%d="Critique"' % (SYN + 6)], fill=fill(RED_BG), font=f(9, True, RED)))
    ws.conditional_formatting.add(
        "E%d:P%d" % (SYN + 6, SYN + 6),
        FormulaRule(formula=['E%d="Vigilance"' % (SYN + 6)], fill=fill(AMBER_BG), font=f(9, True, AMBER)))
    ws.conditional_formatting.add(
        "E%d:P%d" % (SYN + 6, SYN + 6),
        FormulaRule(formula=['E%d="Couvert"' % (SYN + 6)], fill=fill(GREEN_BG), font=f(9, True, GREEN)))
    ws.conditional_formatting.add(
        "E%d:P%d" % (SYN + 5, SYN + 5),
        FormulaRule(formula=['E%d="OUI"' % (SYN + 5)], fill=fill(RED_BG), font=f(9, True, RED)))

    # ---------------------------------------------------------- graphique
    ch = BarChart()
    ch.type = "col"
    ch.style = 2
    ch.title = "Effectif autonome par poste (niveau ≥ 3) et cible"
    ch.height, ch.width = 8.5, 26
    dat = Reference(ws, min_col=P1, max_col=P2, min_row=SYN, max_row=SYN)
    cible = Reference(ws, min_col=P1, max_col=P2, min_row=SYN + 2, max_row=SYN + 2)
    ch.add_data(dat, from_rows=True, titles_from_data=False)
    ch.add_data(cible, from_rows=True, titles_from_data=False)
    ch.set_categories(Reference(ws, min_col=P1, max_col=P2, min_row=9, max_row=9))
    ch.series[0].tx = SeriesLabel(v="Effectif autonome (≥ 3)")
    ch.series[1].tx = SeriesLabel(v="Cible")
    ch.series[0].graphicalProperties = GraphicalProperties(solidFill="1F6FB2")
    ch.series[1].graphicalProperties = GraphicalProperties(solidFill="C9D6E2")
    ch.y_axis.title = "Nombre de personnes"
    ch.y_axis.majorGridlines = None
    ch.gapWidth = 60
    ch.overlap = -20
    str_categories(ch, "'POLYVALENCE'!$E$9:$P$9")
    ws.add_chart(ch, "A%d" % (SYN + 9))

    widths(ws, {"A": 11, "B": 22, "C": 20, "D": 10, "Q": 11, "R": 11, "S": 13, "T": 32})
    for col in range(P1, P2 + 1):
        ws.column_dimensions[gcl(col)].width = 8.5
    ws.freeze_panes = "E12"
    page(ws, "landscape", title_rows="9:11")
    return ws
