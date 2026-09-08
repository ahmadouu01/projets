# -*- coding: utf-8 -*-
"""Onglet PRISE_DE_POSTE : standard de prise de poste, passation et escalade."""
import datetime as dt
from openpyxl.styles import Alignment, Border, Side
from openpyxl.worksheet.datavalidation import DataValidation
from openpyxl.formatting.rule import FormulaRule
from common import *
import data as D

FIRST = 8
LAST = FIRST + len(D.POSTE_CHECK) - 1


def build(wb):
    ws = wb.create_sheet("PRISE_DE_POSTE")
    ws.sheet_view.showGridLines = False
    ws.sheet_properties.tabColor = "5B4B7A"
    NC = 10

    title_band(ws, 1, NC, "STANDARD DE PRISE DE POSTE ET DE PASSATION",
               "Document à imprimer et à tenir à chaque changement d'équipe. Les points bloquants doivent être soldés avant le lancement de la série. "
               "Structure SQCDP : Sécurité, Qualité, Coût, Délai, Personnel.")

    band(ws, 4, 1, NC, "IDENTIFICATION DU POSTE")
    ident = [("A5", "Date :", "B5", dt.date(2026, 1, 30)),
             ("C5", "Ligne / zone :", "D5", "Ligne 1 - Assemblage"),
             ("F5", "Équipe sortante :", "G5", "Équipe A (matin)"),
             ("I5", "Superviseur sortant :", "J5", "DURAND M."),
             ("A6", "Heure :", "B6", dt.time(13, 30)),
             ("C6", "Effectif présent / prévu :", "D6", "11 / 12"),
             ("F6", "Équipe entrante :", "G6", "Équipe B (après-midi)"),
             ("I6", "Superviseur entrant :", "J6", "MARCHAND J.")]
    for lab_ref, lab, val_ref, val in ident:
        label(ws, lab_ref, lab, bold=True, size=9, align="right")
        c = ws[val_ref]
        c.value = val
        c.font = f(9, True, "0000FF")
        c.fill = fill(INPUT_BG)
        c.border = BOX
        c.alignment = Alignment(horizontal="left", vertical="center", indent=1)
        if isinstance(val, dt.date):
            c.number_format = DATE
        if isinstance(val, dt.time):
            c.number_format = HOUR
    for rng in ("D5:E5", "G5:H5", "J5:J5", "D6:E6", "G6:H6"):
        if ":" in rng and rng.split(":")[0] != rng.split(":")[1]:
            ws.merge_cells(rng)
        for row in ws[rng]:
            for cc in row:
                cc.fill = fill(INPUT_BG)
                cc.border = BOX

    header_row(ws, FIRST - 1, ["Séquence", "Jalon", "Axe\nSQCDP", "Point de contrôle",
                               "Bloquant", "Statut", "Constat / observation",
                               "Action immédiate", "Pilote", "Délai"], height=34)

    for i, (seq, jalon, axe, point, bloq) in enumerate(D.POSTE_CHECK):
        r = FIRST + i
        ws["A%d" % r] = seq
        ws["B%d" % r] = jalon
        ws["C%d" % r] = axe
        ws["D%d" % r] = point
        ws["E%d" % r] = bloq
        for col in range(1, 11):
            c = ws.cell(row=r, column=col)
            c.border = BOX
            c.font = f(9)
            c.alignment = Alignment(horizontal="left", vertical="center", indent=1, wrap_text=True)
        for col in (2, 3, 5, 6):
            ws.cell(row=r, column=col).alignment = Alignment(horizontal="center", vertical="center")
        ws.cell(row=r, column=1).font = f(8.5, True, STEEL)
        for col in (6, 7, 8, 9, 10):
            ws.cell(row=r, column=col).font = f(9, color="0000FF")
            ws.cell(row=r, column=col).fill = fill(INPUT_BG)
        ws.cell(row=r, column=10).number_format = DATE
        ws.row_dimensions[r].height = 20

    # bilan
    r = LAST + 2
    band(ws, r - 1, 1, NC, "BILAN DE LA PRISE DE POSTE")
    bilan = [
        ("Taux de conformité du standard",
         '=IFERROR(COUNTIF($F${f}:$F${l},"OK")/(COUNTIF($F${f}:$F${l},"OK")+COUNTIF($F${f}:$F${l},"NOK")),"")'.format(f=FIRST, l=LAST), PCT),
        ("Points de contrôle non conformes", '=COUNTIF($F${f}:$F${l},"NOK")'.format(f=FIRST, l=LAST), NUM),
        ("Dont points bloquants non soldés",
         '=COUNTIFS($E${f}:$E${l},"Oui",$F${f}:$F${l},"NOK")'.format(f=FIRST, l=LAST), NUM),
        ("Décision de lancement",
         ('=IF(COUNTIF($F${f}:$F${l},"OK")+COUNTIF($F${f}:$F${l},"NOK")+COUNTIF($F${f}:$F${l},"S.O.")=0,'
          '"Contrôles non réalisés",'
          'IF(COUNTIFS($E${f}:$E${l},"Oui",$F${f}:$F${l},"NOK")>0,'
          '"LANCEMENT INTERDIT — solder les points bloquants","Lancement autorisé"))').format(f=FIRST, l=LAST), None),
    ]
    for i, (lab, formule, fmt) in enumerate(bilan):
        rr = r + i
        ws.merge_cells(start_row=rr, start_column=1, end_row=rr, end_column=4)
        c = ws.cell(row=rr, column=1, value=lab)
        c.font = f(9.5, True, STEEL)
        c.alignment = Alignment(horizontal="right", vertical="center")
        ws.merge_cells(start_row=rr, start_column=5, end_row=rr, end_column=10)
        v = ws.cell(row=rr, column=5, value=formule)
        v.font = f(11, True, NAVY)
        v.alignment = Alignment(horizontal="left", vertical="center", indent=1)
        if fmt:
            v.number_format = fmt
        for col in range(1, 11):
            ws.cell(row=rr, column=col).border = BOX
            if col >= 5:
                ws.cell(row=rr, column=col).fill = fill(CALC_BG)
            else:
                ws.cell(row=rr, column=col).fill = fill(LIGHT)
        ws.row_dimensions[rr].height = 19
    ws.conditional_formatting.add(
        "E%d:J%d" % (r + 3, r + 3),
        FormulaRule(formula=['LEFT($E%d,9)="LANCEMENT"' % (r + 3)],
                    fill=fill(RED_BG), font=f(11, True, RED)))
    ws.conditional_formatting.add(
        "E%d:J%d" % (r + 3, r + 3),
        FormulaRule(formula=['$E%d="Lancement autorisé"' % (r + 3)],
                    fill=fill(GREEN_BG), font=f(11, True, GREEN)))
    ws.conditional_formatting.add(
        "E%d:J%d" % (r + 3, r + 3),
        FormulaRule(formula=['$E%d="Contrôles non réalisés"' % (r + 3)],
                    fill=fill(LIGHT), font=f(11, True, GREY)))

    # passation ecrite
    r2 = r + 6
    band(ws, r2, 1, NC, "PASSATION ÉCRITE À L'ÉQUIPE SUIVANTE")
    passation = [
        "En-cours, état des lignes et avancement des OF",
        "Aléas non résolus et actions engagées",
        "Consignes particulières (client, qualité, sécurité, dérogation)",
        "Points de vigilance pour le poste suivant",
        "Interventions de maintenance en cours ou attendues",
    ]
    for i, lab in enumerate(passation):
        rr = r2 + 1 + i
        ws.merge_cells(start_row=rr, start_column=1, end_row=rr, end_column=3)
        c = ws.cell(row=rr, column=1, value=lab)
        c.font = f(9, True, STEEL)
        c.alignment = Alignment(horizontal="left", vertical="center", indent=1, wrap_text=True)
        ws.merge_cells(start_row=rr, start_column=4, end_row=rr, end_column=10)
        v = ws.cell(row=rr, column=4)
        v.font = f(9, color="0000FF")
        v.alignment = Alignment(horizontal="left", vertical="center", indent=1, wrap_text=True)
        for col in range(1, 11):
            ws.cell(row=rr, column=col).border = BOX
            ws.cell(row=rr, column=col).fill = fill(LIGHT if col <= 3 else INPUT_BG)
        ws.row_dimensions[rr].height = 30

    # matrice d'escalade
    r3 = r2 + 7 + len(passation) - 5
    r3 = r2 + len(passation) + 3
    band(ws, r3, 1, NC, "MATRICE D'ESCALADE  —  qui alerter, dans quel délai")
    header_row(ws, r3 + 1, ["Niveau", "Situation déclenchante", "Délai de réaction",
                            "Interlocuteur", "Moyen d'alerte", "Trace obligatoire",
                            "", "", "", ""], height=28)
    for i, ligne in enumerate(D.ESCALADE):
        rr = r3 + 2 + i
        for j, v in enumerate(ligne):
            c = ws.cell(row=rr, column=1 + j, value=v)
            c.border = BOX
            c.font = f(9, j == 0)
            c.alignment = Alignment(horizontal="left" if j else "center",
                                    vertical="center", indent=1, wrap_text=True)
        for col in range(7, 11):
            ws.cell(row=rr, column=col).border = BOX
        ws.row_dimensions[rr].height = 26
    ws.merge_cells(start_row=r3 + 1, start_column=6, end_row=r3 + 1, end_column=10)
    for i in range(len(D.ESCALADE)):
        ws.merge_cells(start_row=r3 + 2 + i, start_column=6, end_row=r3 + 2 + i, end_column=10)

    # validation
    r4 = r3 + 3 + len(D.ESCALADE)
    band(ws, r4, 1, NC, "VISA DE PASSATION")
    for i, lab in enumerate(["Superviseur sortant (nom et visa)", "Superviseur entrant (nom et visa)"]):
        rr = r4 + 1
        c1 = 1 + i * 5
        ws.merge_cells(start_row=rr, start_column=c1, end_row=rr, end_column=c1 + 1)
        c = ws.cell(row=rr, column=c1, value=lab)
        c.font = f(9, True, STEEL)
        c.alignment = Alignment(horizontal="right", vertical="center")
        ws.merge_cells(start_row=rr, start_column=c1 + 2, end_row=rr, end_column=c1 + 4)
        for col in range(c1, c1 + 5):
            ws.cell(row=rr, column=col).border = BOX
            ws.cell(row=rr, column=col).fill = fill(INPUT_BG if col >= c1 + 2 else LIGHT)
        ws.row_dimensions[rr].height = 30

    dv = DataValidation(type="list", formula1="LST_CONTROLE", allow_blank=True,
                        showErrorMessage=True, errorTitle="Statut",
                        error="Choisir OK, NOK ou S.O.")
    ws.add_data_validation(dv)
    dv.add("F%d:F%d" % (FIRST, LAST))
    dv2 = DataValidation(type="list", formula1="LST_PILOTES", allow_blank=True)
    ws.add_data_validation(dv2)
    dv2.add("I%d:I%d" % (FIRST, LAST))

    ws.conditional_formatting.add(
        "A%d:J%d" % (FIRST, LAST),
        FormulaRule(formula=['AND($F%d="NOK",$E%d="Oui")' % (FIRST, FIRST)],
                    fill=fill(RED_BG), font=f(9, True, RED)))
    ws.conditional_formatting.add(
        "F%d:F%d" % (FIRST, LAST),
        FormulaRule(formula=['$F%d="OK"' % FIRST], fill=fill(GREEN_BG), font=f(9, True, GREEN)))
    ws.conditional_formatting.add(
        "F%d:F%d" % (FIRST, LAST),
        FormulaRule(formula=['$F%d="NOK"' % FIRST], fill=fill(RED_BG), font=f(9, True, RED)))
    ws.conditional_formatting.add(
        "E%d:E%d" % (FIRST, LAST),
        FormulaRule(formula=['$E%d="Oui"' % FIRST], font=f(9, True, RED)))

    widths(ws, {"A": 24, "B": 14, "C": 11, "D": 58, "E": 9, "F": 9,
                "G": 30, "H": 30, "I": 14, "J": 12})
    ws.freeze_panes = "A8"
    page(ws, "landscape", title_rows="7:7")
    return ws
